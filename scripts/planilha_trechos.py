"""Planilha Excel com dinâmicas a partir dos JSON canônicos do site (dados/alimentadores/*.json).

python planilha_trechos.py <pasta_json> <manifesto.json> <saida.xlsx>

Abas: Leia-me, Dinâmica km, Dinâmica NAE, Trechos (1 linha por trecho × município), NAE (1 linha por ocorrência).
As dinâmicas são criadas pelo LibreOffice (UNO) e se atualizam ao abrir no Excel.
"""
import glob, json, os, re, subprocess, sys, time, zipfile, shutil, tempfile  # noqa
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

MES = dict(zip(['01', '02', '03', '04', '05', '06', '07', '08', '09', '10', '11', '12'],
               ['jan', 'fev', 'mar', 'abr', 'mai', 'jun', 'jul', 'ago', 'set', 'out', 'nov', 'dez']))
CAB_T = ['Polo', 'SE', 'Subestação', 'Município', 'Alimentador', 'Trecho', 'Classe', 'Tipo', 'Ativos', 'Caminho', 'Chave início',
         'Tipo chave início', 'Chave(s) fim', 'Fase', 'Km no município', 'Km do trecho (total)', 'NAE no trecho', 'Postes']
LIM = 32000                                     # limite de texto de uma célula do Excel (32.767)


def ativos_por_trecho(d):
    """(trecho, município) → chaves e postes que ficam nele, na ordem da rede a partir do DJ.

    Chave: no trecho que ela abre (início); senão no trecho que passa por ela; senão no trecho dos vãos do poste
    dela (o que não termina nela). Município: o do vão que chega no poste. Postes N123 (sem número) ficam de fora.
    """
    ts = d['trechos']
    no_poste = {}                                 # poste → [(ordem do vão, trecho, município)]
    for i, v in enumerate(d['vaos']):
        for p in (v[6], v[7]):
            if p:
                no_poste.setdefault(p, []).append((i, ts[v[4]]['nome'], v[8]))
    ch, po = {}, {}
    for p, lst in no_poste.items():
        if not re.fullmatch(r'N\d+', p):
            for i, t, mun in lst:
                po.setdefault((t, mun), {}).setdefault(p, i)
    inicio = {t['inicio']['chave']: t['nome'] for t in ts}
    passa = {c: t['nome'] for t in ts for c in (t.get('passa') or [])}
    fim = {}
    for t in ts:
        for c in t.get('fim') or []:
            fim.setdefault(c, set()).add(t['nome'])
    for c in d['chaves']:
        n, lst = c['nome'], no_poste.get(c.get('poste'), [])
        alvo = inicio.get(n) or passa.get(n)
        if not alvo and lst:
            fora_fim = [t for _, t, _ in lst if t not in fim.get(n, ())]
            alvo = fora_fim[0] if fora_fim else lst[0][1]
        if not alvo:
            continue
        dele = [(i, mun) for i, t, mun in lst if t == alvo] or [(0, None)]
        ordem = -1 if inicio.get(n) == alvo else dele[0][0]
        ch.setdefault((alvo, dele[0][1]), []).append((ordem, n))
    junta = lambda xs: (lambda s: s if len(s) <= LIM else s[:LIM].rsplit('; ', 1)[0] + f'; … (+{len(xs)} no total)')('; '.join(xs))
    return ({k: junta([n for _, n in sorted(v)]) for k, v in ch.items()},
            {k: junta([p for p, _ in sorted(v.items(), key=lambda x: x[1])]) for k, v in po.items()})
CAB_N = ['Mês', 'Data/hora', 'Polo', 'SE', 'Município (do trecho)', 'Alimentador', 'Trecho', 'Classe', 'Ativo', 'Abrangência',
         'Causa', 'Subcausa', 'Programada', 'Consumidores', 'Duração (min)', 'CHI (h)', 'Regra de associação', 'Distância (m)']


def linhas(pasta, nomes):
    tr, oc, fora, meses, crit = [], [], [], set(), set()
    for f in sorted(glob.glob(os.path.join(pasta, '*.json'))):
        d = json.load(open(f, encoding='utf-8'))
        if d.get('validacao', {}).get('status') == 'falha':
            fora.append(d['alimentador']); continue
        al, polo, se, sen = d['alimentador'], d.get('polo'), d.get('se'), d.get('se_nome')
        meses.update(d['nae'].get('meses', [])); crit.update((d.get('origem') or {}).get('critica') or [])
        principal, classe, nae = {}, {}, {}
        for t in d['trechos']:
            mm = t.get('municipios') or {}
            principal[t['nome']] = max(mm, key=mm.get) if mm else None
            classe[t['nome']] = t['classe']
        for o in d['nae']['oc']:
            if o.get('trecho'): nae[o['trecho']] = nae.get(o['trecho'], 0) + 1
        ats, pos = ativos_por_trecho(d)
        for t in d['trechos']:
            mm = t.get('municipios') or {None: t['ext_m']}
            for k, cod in enumerate(c for c, _ in sorted(mm.items(), key=lambda x: -x[1])):
                m = mm[cod]
                # chave sem município (fora da malha) vai para a 1ª linha do trecho
                a = '; '.join(x for x in (ats.get((t['nome'], cod)), ats.get((t['nome'], None)) if k == 0 else None) if x)
                tr.append([polo, se, sen, nomes.get(cod, cod or '(fora da malha)'), al, t['nome'], t['classe'], t.get('tipo'), a,
                           ' › '.join(t['caminho']), t['inicio']['chave'], t['inicio']['tipo'],
                           ', '.join(t.get('fim') or []) or 'fim de linha', t.get('fase') or '', round(m / 1000, 3),
                           round(t['ext_m'] / 1000, 3), nae.get(t['nome'], 0) if cod == principal[t['nome']] else 0,
                           pos.get((t['nome'], cod), '')])
        for o in d['nae']['oc']:
            tn = o.get('trecho')
            y, mo = o['mes'].split('-')
            oc.append([f'{mo}-{MES[mo]}/{y}', o.get('dia'), polo, se,
                       nomes.get(principal.get(tn), '(sem trecho)') if tn else '(sem trecho)', al, tn or '(sem trecho)',
                       classe.get(tn, '(sem trecho)'), o.get('ativo'), o.get('abr'), o.get('causa'), o.get('sub'),
                       'Sim' if o.get('prog') else 'Não', o.get('cons') or 0, o.get('dur') or 0, None, o.get('regra'), o.get('dist_m')])
    return tr, oc, fora, sorted(meses), sorted(crit)


def base(saida, tr, oc, fora, meses, crit, n_al):
    F, FB = Font(name='Arial', size=10), Font(name='Arial', size=10, bold=True, color='FFFFFF')
    wb = Workbook(); ws = wb.active; ws.title = 'Leia-me'
    per = f"{MES[meses[0][5:]]}–{MES[meses[-1][5:]]}/{meses[-1][:4]}" if meses else '—'
    sem = sum(1 for r in oc if r[6] == '(sem trecho)')
    txt = [f'Trechos por alimentador — {n_al} alimentadores publicados',
           f'Fonte: JSON canônico do site (dados/alimentadores). NAE: Crítica {per}, só Meio Ambiente · Árvore/Eucalipto.',
           f"Arquivos da Crítica: {', '.join(crit) or '—'}", '',
           'Abas:',
           f'• Trechos — {len(tr):,} linhas, 1 por trecho × município. "Km no município" soma certo em qualquer dinâmica.'.replace(',', '.'),
           '  "Ativos": chaves do trecho (DJ, religador, fusível, faca, trip saver…) separadas por "; ", na ordem a partir do DJ;'
           ' "Postes": postes numerados do trecho. Os dois por município (cada linha traz os que ficam nela).',
           f'• NAE — {len(oc):,} ocorrências, 1 por linha. Conte linhas para ter o NAE; some Consumidores/CHI.'.replace(',', '.'),
           '• Dinâmica km — extensão por município → alimentador × classe (T1/T2/T3).',
           '• Dinâmica NAE — ocorrências por município → alimentador → classe × mês.',
           '• Ranking trechos — trechos com NAE, do que tem mais para o que tem menos (NAE, consumidores e CHI por fórmula da aba NAE).', '',
           'Classes: T1 = DJ → 1º religador; T2 = entre religadores (T2-A, T2-B… por religador); T3-A… = depois de cada chave fusível; sem T4.', '',
           'Premissas:',
           '• Trecho em mais de um município: o km é dividido pela rede de cada município (amostra a cada 20 m).',
           '• NAE do trecho vai para o município onde o trecho tem mais rede (a Crítica não traz coordenada da ocorrência).',
           '  Na aba Trechos, "NAE no trecho" fica só na linha desse município (nas outras = 0), para não contar duas vezes.',
           '• CHI = Consumidores × Duração (min) ÷ 60 (fórmula na aba NAE).',
           f'• "(sem trecho)": {sem:,} ocorrências cujo ativo não foi achado na rede; entram no total do alimentador, sem classe.'.replace(',', '.'),
           f"• Fora da planilha (bloqueados na validação do site): {', '.join(fora) or 'nenhum'}.", '',
           'Atualizar: Dados → Atualizar Tudo (as dinâmicas também se atualizam ao abrir).']
    for i, s in enumerate(txt, 1):
        c = ws.cell(i, 1, s); c.font = Font(name='Arial', size=12 if i == 1 else 10, bold=i in (1, 5, 13))
    ws.column_dimensions['A'].width = 130

    def tabela(nome, cab, rows, larg, fmt):
        w = wb.create_sheet(nome); w.append(cab)
        for r in rows: w.append(r)
        n = len(rows) + 1
        for c in w[1]:
            c.font = FB; c.fill = PatternFill('solid', fgColor='1F4E78'); c.alignment = Alignment(wrap_text=True, vertical='center')
        for i, l in enumerate(larg): w.column_dimensions[get_column_letter(i + 1)].width = l
        for row in w.iter_rows(min_row=2):
            for c in row: c.font = F
        for col, f in fmt.items():
            for r in range(2, n + 1): w[f'{col}{r}'].number_format = f
        w.freeze_panes = 'A2'
        t = Table(displayName='tb' + nome, ref=f'A1:{get_column_letter(len(cab))}{n}')
        t.tableStyleInfo = TableStyleInfo(name='TableStyleLight9', showRowStripes=True); w.add_table(t)
        return w, n
    tabela('Trechos', CAB_T, tr, [8, 10, 22, 18, 13, 9, 7, 10, 40, 30, 14, 9, 22, 11, 12, 12, 10, 40], {'O': '#,##0.000', 'P': '#,##0.000'})
    w, n = tabela('NAE', CAB_N, oc, [11, 12, 8, 10, 18, 13, 12, 11, 13, 10, 14, 40, 10, 12, 11, 10, 24, 10],
                  {'N': '#,##0', 'O': '#,##0', 'P': '#,##0.0'})
    for r in range(2, n + 1): w[f'P{r}'] = f'=N{r}*O{r}/60'
    # ranking por trecho: ordem calculada aqui; os números vêm por fórmula da aba NAE
    info = {}
    for r in tr:
        info.setdefault((r[4], r[5]), (r[0], r[3], r[6], r[15]))   # polo, município da 1ª linha (a de mais rede), classe, km
    soma = {}
    for r in oc:
        if r[6] != '(sem trecho)':
            x = soma.setdefault((r[5], r[6]), [0, 0, 0]); x[0] += 1; x[1] += r[13]; x[2] += r[13] * r[14]
    ordem = sorted(soma, key=lambda k: (-soma[k][0], -soma[k][2], k))
    cab = ['Posição', 'Polo', 'Município', 'Alimentador', 'Trecho', 'Classe', 'Km do trecho', 'NAE', 'Consumidores', 'CHI (h)', 'NAE / 100 km']
    linhas_r = [[i + 1, *info[k][:2], k[0], k[1], info[k][2], info[k][3], None, None, None, None] for i, k in enumerate(ordem)]
    w, n = tabela('RankingTrechos', cab, linhas_r, [8, 8, 20, 13, 9, 7, 11, 8, 13, 10, 11],
                  {'G': '#,##0.000', 'I': '#,##0', 'J': '#,##0.0', 'K': '#,##0.0'})
    w.title = 'Ranking trechos'
    for r in range(2, n + 1):
        cond = f'NAE!$F:$F,$D{r},NAE!$G:$G,$E{r}'
        w[f'H{r}'] = f'=COUNTIFS({cond})'
        w[f'I{r}'] = f'=SUMIFS(NAE!$N:$N,{cond})'
        w[f'J{r}'] = f'=SUMIFS(NAE!$P:$P,{cond})'
        w[f'K{r}'] = f'=IF(G{r}>0,H{r}/G{r}*100,0)'
    wb.create_sheet('Dinâmica km', 1); wb.create_sheet('Dinâmica NAE', 2)
    wb.save(saida)


def dinamicas(entrada, saida):
    """Abre no LibreOffice sem tela, cria as duas tabelas dinâmicas e salva como xlsx."""
    porta = '2002'
    perfil = tempfile.mkdtemp()
    proc = subprocess.Popen(['soffice', '--headless', '--invisible', '--norestore', f'-env:UserInstallation=file://{perfil}',
                             f'--accept=socket,host=localhost,port={porta};urp;'], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    try:
        import uno
        from com.sun.star.beans import PropertyValue
        from com.sun.star.table import CellAddress
        from com.sun.star.sheet.DataPilotFieldOrientation import ROW, COLUMN, DATA, PAGE
        from com.sun.star.sheet.GeneralFunction import SUM, COUNT

        def P(n, v):
            p = PropertyValue(); p.Name = n; p.Value = v; return p
        ctx = uno.getComponentContext()
        res = ctx.ServiceManager.createInstanceWithContext('com.sun.star.bridge.UnoUrlResolver', ctx)
        for _ in range(60):
            try:
                c = res.resolve(f'uno:socket,host=localhost,port={porta};urp;StarOffice.ComponentContext'); break
            except Exception:
                time.sleep(1)
        desk = c.ServiceManager.createInstanceWithContext('com.sun.star.frame.Desktop', c)
        doc = desk.loadComponentFromURL(uno.systemPathToFileUrl(os.path.abspath(entrada)), '_blank', 0, (P('Hidden', True),))
        sh = doc.Sheets
        nomes = [sh.getByIndex(i).Name for i in range(sh.Count)]

        def piv(dst, src, nome, rows, cols, dado, pages):
            s = sh.getByName(src); cur = s.createCursor(); cur.gotoEndOfUsedArea(False)
            ra = s.getCellRangeByPosition(0, 0, cur.RangeAddress.EndColumn, cur.RangeAddress.EndRow).RangeAddress
            tabs = sh.getByName(dst).getDataPilotTables(); desc = tabs.createDataPilotDescriptor(); desc.setSourceRange(ra)
            f = desc.getDataPilotFields()
            fld = {f.getByIndex(i).Name: f.getByIndex(i) for i in range(f.getCount())}
            for n in rows: fld[n].Orientation = ROW
            for n in cols: fld[n].Orientation = COLUMN
            for n in pages: fld[n].Orientation = PAGE
            fld[dado[0]].Orientation = DATA; fld[dado[0]].Function = dado[1]
            a = CellAddress(); a.Sheet = nomes.index(dst); a.Column = 0; a.Row = len(pages) + 2
            tabs.insertNewByName(nome, a, desc)
        piv('Dinâmica km', 'Trechos', 'dinKm', ['Município', 'Alimentador'], ['Classe'], ('Km no município', SUM), ['Polo', 'Subestação', 'Fase'])
        piv('Dinâmica NAE', 'NAE', 'dinNae', ['Município (do trecho)', 'Alimentador', 'Classe'], ['Mês'], ('Ativo', COUNT), ['Polo', 'SE', 'Programada'])
        doc.storeToURL(uno.systemPathToFileUrl(os.path.abspath(saida)), (P('FilterName', 'Calc MS Excel 2007 XML'),))
        doc.close(True)
    finally:
        proc.terminate(); proc.wait(30); shutil.rmtree(perfil, ignore_errors=True)


def ajusta(arq):
    """Atualizar ao abrir, nomes dos valores e rótulos em português."""
    tmp = arq + '.tmp'
    with zipfile.ZipFile(arq) as zi, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            b = zi.read(it.filename)
            if it.filename.startswith('xl/pivotCache/pivotCacheDefinition'):
                b = b.replace(b'<pivotCacheDefinition ', b'<pivotCacheDefinition refreshOnLoad="1" ', 1)
            elif it.filename.startswith('xl/pivotTables/pivotTable'):
                s = b.decode('utf-8')
                s = re.sub(r'<dataField name="Sum - [^"]*"', '<dataField name="Km"', s)
                s = re.sub(r'<dataField name="Count - [^"]*"', '<dataField name="NAE (ocorrências)"', s)
                b = s.encode('utf-8')
            elif it.filename == 'xl/sharedStrings.xml':
                s = b.decode('utf-8')
                for a, c in [('>Filter<', '>Filtros<'), ('>- all -<', '>(Tudo)<'), ('>Total Result<', '>Total Geral<')]:
                    s = s.replace(a, c)
                s = re.sub(r'>Sum - Km no município<', '>Km<', s); s = re.sub(r'>Count - Ativo<', '>NAE (ocorrências)<', s)
                b = s.encode('utf-8')
            zo.writestr(it, b)
    os.replace(tmp, arq)


def main(pasta, manifesto, saida):
    man = json.load(open(manifesto, encoding='utf-8'))
    nomes = {m['cod']: m['nome'] for m in man['municipios']}
    tr, oc, fora, meses, crit = linhas(pasta, nomes)
    n_al = len({r[4] for r in tr})
    b = saida + '.base.xlsx'
    base(b, tr, oc, fora, meses, crit, n_al)
    dinamicas(b, saida); os.remove(b); ajusta(saida)
    print(f'{n_al} alimentadores · {len(tr)} linhas de trecho · {len(oc)} ocorrências · bloqueados: {len(fora)} → {saida}')


if __name__ == '__main__':
    main(*sys.argv[1:4])
