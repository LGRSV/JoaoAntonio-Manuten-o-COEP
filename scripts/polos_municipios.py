"""Município → polo DCMD e regional, a partir da planilha de cidades por polo.

python scripts/polos_municipios.py Cidades_por_Polo_e_Regional_DCMD_MAIUSCULAS_SEM_ACENTO.xlsx

Grava alimentador/site/polos.json com o código IBGE de cada município (casado pelo nome com a malha do mapa da
página inicial). Distritos que não são município (Vila Luzimangues, Taquaralto) ficam de fora: a rede deles conta
no município a que pertencem.
"""
import json
import re
import sys
import unicodedata
from pathlib import Path

from openpyxl import load_workbook

PKG = Path(__file__).resolve().parent.parent / 'alimentador'
# polo da planilha → nome com acento e equipe de rede de distribuição do polo (ordem: Norte, Centro, Sul)
POLOS = {'ARAGUAINA': ('Araguaína', 'ETO-RD-AR'), 'AUGUSTINOPOLIS': ('Augustinópolis', 'ETO-RD-AG'),
         'GUARAI': ('Guaraí', 'ETO-RD-GR'), 'PALMAS': ('Palmas', 'ETO-RD-PA'), 'PARAISO': ('Paraíso', 'ETO-RD-PS'),
         'PORTO NACIONAL': ('Porto Nacional', 'ETO-RD-PO'), 'GURUPI': ('Gurupi', 'ETO-RD-GU'),
         'DIANOPOLIS': ('Dianópolis', 'ETO-RD-DP')}
# nome da planilha → nome do IBGE, quando diferem
APELIDO = {'DIVINOPOLIS': 'DIVINOPOLIS DO TOCANTINS', 'MARIANOPOLIS': 'MARIANOPOLIS DO TOCANTINS',
           'PONTE ALTA TOCANTINS': 'PONTE ALTA DO TOCANTINS', 'PONTE ALTA BOM JESUS': 'PONTE ALTA DO BOM JESUS'}
DISTRITO = {'VILA LUZIMANGUES': 'Porto Nacional', 'TAQUARALTO': 'Palmas'}


def norm(s):
    s = unicodedata.normalize('NFD', str(s)).encode('ascii', 'ignore').decode().upper().replace("'", '')
    return re.sub(r'\s+', ' ', re.sub(r'[^A-Z ]', ' ', s)).strip()


def main(xlsx):
    geo = {norm(m['nome']): m for m in json.loads((PKG / 'site' / 'mapa-municipios-geo.json').read_text(encoding='utf-8'))['muns']}
    linhas = list(load_workbook(xlsx, read_only=True, data_only=True)['POLO_REGIONAL_CIDADE'].iter_rows(values_only=True))[1:]
    polos = {p: dict(id=norm(p).lower().replace(' ', '-'), nome=n, regional=None, equipe=e, municipios=[])
             for p, (n, e) in POLOS.items()}
    fora = []
    for cidade, polo, regional in linhas:
        c, p = norm(cidade), norm(polo)
        if c in DISTRITO:
            fora.append(f'{cidade.title()} (distrito de {DISTRITO[c]}: a rede conta em {DISTRITO[c]})')
            continue
        m = geo[APELIDO.get(c, c)]                     # KeyError = nome novo na planilha: ajuste APELIDO
        polos[p]['regional'] = regional.title()
        polos[p]['municipios'].append(m['cod'])
    todos = [cod for p in polos.values() for cod in p['municipios']]
    assert len(todos) == len(set(todos)) == len(geo), 'cada município do TO em um polo só'
    for p in polos.values():
        p['municipios'].sort()
    saida = dict(fonte=Path(xlsx).name, polos=list(polos.values()), fora=fora)
    (PKG / 'site' / 'polos.json').write_text(json.dumps(saida, ensure_ascii=False, indent=1), encoding='utf-8')
    print(', '.join(f"{p['nome']} {len(p['municipios'])}" for p in polos.values()), '| fora:', '; '.join(fora))


if __name__ == '__main__':
    main(sys.argv[1])
