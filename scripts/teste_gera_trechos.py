"""Teste mínimo do pipeline por município.  python scripts/teste_gera_trechos.py

Rede inventada (DJ → religador → fusível) cortada ao meio por dois municípios fictícios: confere a divisão por
município (amostra de 20 m), os arquivos gerados, os links, a validação e que duas execuções dão o mesmo resultado.
"""
import hashlib
import json
import math
import sys
import shutil
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import gera_trechos  # noqa: E402
import limpeza_faixa  # noqa: E402
import ativos_gis  # noqa: E402
import atualiza_critica  # noqa: E402
import inspecao  # noqa: E402
import atualiza_bases  # noqa: E402
import manutencao  # noqa: E402
from alimentador.testes import _kml_daimon  # noqa: E402


def rede(d):
    y = lambda i: -7.2 + i * 0.0003                   # linha norte-sul, ~33 m por vão
    P = {i: (y(i), -48.2) for i in range(8)}
    meio = lambda i: ((y(i) + y(i + 1)) / 2, -48.2)
    chaves = [('DJ_ALTESTE_FICT', *meio(-1), False), ('7900001', *meio(2), False), ('0300001', *meio(5), False),
              ('3100001', *meio(4), False)]             # 3100001: ponto sem suporte (fora da topologia), no vão t4
    tr = [('Suporte Chave DJ_ALTESTE_FICT', (y(-1), -48.2), P[0]), ('t1', P[0], P[1]), ('t2', P[1], P[2]),
          ('Suporte Chave 7900001', P[2], P[3]), ('t3', P[3], P[4]), ('t4', P[4], P[5]),
          ('Suporte Chave 0300001', P[5], P[6]), ('t5', P[6], P[7])]
    (d / 'kml').mkdir()
    (d / 'kml' / 'SE-XTST_Circ-ALTESTE.kml').write_text(_kml_daimon(chaves, tr), encoding='utf-8')
    corte = y(3.5)                                    # divisa no meio do vão t3
    caixa = lambda c, n, s, nm: dict(type='Feature', properties=dict(id=c, name=nm), geometry=dict(
        type='Polygon', coordinates=[[[-48.3, s], [-48.1, s], [-48.1, n], [-48.3, n], [-48.3, s]]]))
    malha = dict(type='FeatureCollection', features=[caixa('1', corte, -7.3, 'Sul Teste'), caixa('2', -7.1, corte, 'Norte Teste')])
    (d / 'malha.json').write_text(json.dumps(malha), encoding='utf-8')


def assinatura(pasta):
    return {str(p.relative_to(pasta)): hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(pasta.rglob('*')) if p.is_file()}


def test_pipeline():
    with tempfile.TemporaryDirectory() as d:
        d = Path(d)
        rede(d)
        args = ['--kml', str(d / 'kml'), '--malha', str(d / 'malha.json'), '--saida', str(d / 'site'), '--data', '2026-01-01',
                '--alimentador', 'ALTESTE']
        assert gera_trechos.main(args) == 0
        c = json.loads((d / 'site' / 'dados' / 'alimentadores' / 'ALTESTE.json').read_text(encoding='utf-8'))
        assert [t['nome'] for t in c['trechos']] == ['T1', 'T2'] and c['validacao']['status'] != 'falha'
        assert c['resumo']['fusiveis_tronco'] == 1, 'fusível na linha principal não abre T3'
        muns = {m['nome']: m['m'] for m in c['municipios']}
        assert set(muns) == {'Sul Teste', 'Norte Teste'} and abs(sum(muns.values()) - c['resumo']['extensao_m']) < 1
        assert len(c['limitrofes']) == 1 and c['limitrofes'][0]['trecho'] == 'T2', 'vão da divisa medido por amostra'
        assert set(c['trechos'][1]['municipios']) == {'1', '2'} and c['trechos'][0]['caminho'] == ['T1']
        assert c['trechos'][1]['caminho'] == ['T1', 'T2']
        fora = {p[0]: p for p in c['chaves_sem_topologia']}
        assert fora['3100001'][3:] == ['T2', 16.7], 'chave fora da topologia: trecho e distância da ponta de vão mais perto'
        assert all(-51 < v[0] < -45 and -14 < v[1] < -5 for v in c['vaos']), 'coordenadas [lon, lat]'
        man = json.loads((d / 'site' / 'manifesto.json').read_text(encoding='utf-8'))
        assert [m['slug'] for m in man['municipios']] == ['norte-teste', 'sul-teste']
        ativos = d / 'site' / 'dados' / 'ativos'
        hit = json.loads((ativos / (gera_trechos.fatia('7900001') + '.json')).read_text(encoding='utf-8'))['7900001']
        assert hit[0][:4] == ['ALTESTE', '79', 'T2', 'início do T2 · fim do T1'], 'busca de ativo: o religador abre o T2'
        assert gera_trechos.chave_ativo(' 0300-001 ') == '300001', 'busca ignora zeros à esquerda e pontuação'
        for f in ('index.html', 'municipios/norte-teste/index.html', 'alimentadores/ALTESTE/index.html',
                  'alimentadores/ALTESTE/dados.js', 'alimentadores/ALTESTE/ALTESTE-trechos.kmz', 'relatorio.md',
                  'assets/alimentador.js', 'assets/extras.js', 'assets/leaflet/leaflet.js',
                  'assets/busca-ativo.js', 'polos/index.html', 'assets/ranking-polos.js', 'assets/ranking-polos-dados.js',
                  'assets/croqui.js', 'assets/croqui-modelo.js', 'assets/inspecao.js', 'assets/executadas.js',
                  'alimentadores/ALTESTE/executadas.js'):
            assert (d / 'site' / f).exists(), f
        pag = (d / 'site' / 'alimentadores' / 'ALTESTE' / 'index.html').read_text(encoding='utf-8')
        assert 'cdnjs' not in pag and 'const D' not in pag and len(pag) < 20000, 'página leve: dados e código em arquivos'
        assert 'assets/croqui.js' in pag and 'croqui-modelo.js' not in pag, 'croqui na página; o modelo só carrega quando pedido'
        assert '<script src="inspecao.js"></script>' in pag and 'assets/inspecao.js' in pag, 'inspeção na página'
        antes = assinatura(d / 'site')
        assert gera_trechos.main(args) == 0
        assert assinatura(d / 'site') == antes, 'segunda execução igual à primeira'
        assert gera_trechos.main(['--saida', str(d / 'site'), '--validar']) == 0
        assert gera_trechos.main(['--saida', str(d / 'site'), '--so-interface']) == 0
        assert assinatura(d / 'site') == antes, 'só a interface: mesmo resultado'
        divisas_no_site(d, antes)
        ativos_gis_no_site(d)
        critica_sem_kml(d)
        inspecao_no_site(d)
        limpeza_reaproveita(d)
        bases_pelo_menu(d)
        manutencao_no_site(d)
        executadas_no_site(d)


def divisas_no_site(d, antes):
    """divisas.js do alimentador: os dois municípios da malha de teste, anéis fechados recortados em volta da rede, a
    divisa em comum com os mesmos pontos nos dois; sem a malha fica o que já existe (ou null); com ela, igual de novo."""
    site, arq = d / 'site', d / 'site' / 'alimentadores' / 'ALTESTE' / 'divisas.js'
    mu = gera_trechos.le_js(arq)
    assert [m['nome'] for m in mu] == ['Norte Teste', 'Sul Teste'] or [m['nome'] for m in mu] == ['Sul Teste', 'Norte Teste'], mu
    c = json.loads((site / 'dados' / 'alimentadores' / 'ALTESTE.json').read_text(encoding='utf-8'))
    assert [m['cod'] for m in mu] == [m['cod'] for m in c['municipios']], 'na ordem dos municípios da rede'
    (s_, w), (n, e) = mu[0]['corte']
    assert all(s_ < v[1] < n and w < v[0] < e for v in c['vaos'] for v in (v[:2], v[2:4])), 'a rede dentro do corte'
    for m in mu:
        assert m['corte'] == mu[0]['corte'] and m['linhas']
        for r in m['linhas']:
            assert len(r) >= 4 and r[0] == r[-1], 'anel fechado'
            assert all(s_ <= la <= n and w <= lo <= e and la == round(la, 5) and lo == round(lo, 5) for la, lo in r)
    lados = [{tuple(sorted((tuple(p), tuple(q)))) for r in m['linhas'] for p, q in zip(r, r[1:])} for m in mu]
    corte = round(-7.2 + 3.5 * 0.0003, 5)
    comum = lados[0] & lados[1]
    assert comum and all(p[0] == q[0] == corte for p, q in comum), 'divisa em comum igual nos dois (a página desenha uma vez)'
    pag = (site / 'alimentadores' / 'ALTESTE' / 'index.html').read_text(encoding='utf-8')
    assert ('assets/extras.js"></script>\n<script src="divisas.js"></script>\n<script src="../../dados/tocantins.js"></script>\n'
            '<script src="../../assets/municipios.js">' in pag) and (site / 'assets' / 'municipios.js').exists() \
        and (site / 'dados' / 'tocantins.js').exists(), 'divisas e contorno do TO logo depois do extras.js'
    vazio = d / 'cache-vazio'
    assert gera_trechos.main(['--saida', str(site), '--so-interface', '--cache', str(vazio)]) == 0
    assert assinatura(site) == antes and not vazio.exists(), 'sem malha: fica o divisas.js que já existe, sem baixar'
    arq.unlink()
    assert gera_trechos.main(['--saida', str(site), '--so-interface', '--cache', str(vazio)]) == 0
    assert arq.read_text(encoding='utf-8') == 'window.MUNICIPIOS = null;\n', 'sem malha e sem arquivo: null'
    assert gera_trechos.main(['--saida', str(site), '--so-interface', '--malha', str(d / 'malha.json')]) == 0
    assert assinatura(site) == antes, 'só a interface com a malha: o mesmo divisas.js'
    assert all('uf' not in m for m in mu), 'divisa do TO longe da rede: não vem'
    # divisa do Tocantins a ~1 km a oeste da rede: lados que só um município usa, recortados na caixa
    caixa = lambda c_, n_, s_: (c_, c_, [[[[-48.21, s_], [-48.1, s_], [-48.1, n_], [-48.21, n_], [-48.21, s_]]]])
    malha = gera_trechos.Malha([caixa('1', corte, -7.3), caixa('2', -7.1, corte)], 'teste')
    mu = json.loads(gera_trechos.divisas_js(malha, c)[len('window.MUNICIPIOS = '):-2])
    for m in mu:
        assert len(m['uf']) == 1, m
        a_, de, ate = m['uf'][0]
        lado = m['linhas'][a_][de:ate + 1]
        assert len(lado) == 2 and all(lo == -48.21 for la, lo in lado), 'divisa do TO: o lado oeste, até o corte e a divisa'
    assert {tuple(p) for m in mu for a_, de, ate in m['uf'] for p in m['linhas'][a_][de:ate + 1]} == \
        {(s_, -48.21), (corte, -48.21), (n, -48.21)}
    print('ok divisas_no_site')


def limpeza_reaproveita(d):
    """Limpeza sem KMZ: OS já publicada mantém a posição e troca situação/data; OS nova pela chave do mapa; OS que saiu
    da planilha sai do site; executada de equipe de fora não entra."""
    import datetime as dt
    import openpyxl
    site = d / 'site'
    dados = inspecao.le_js(site / 'alimentadores' / 'ALTESTE' / 'dados.js')
    ids = [v[5] for v in dados['vaos']][:2]
    velha = dict(os='000001/2025', os_completa='ETO-RD-PA 000001/2025', sit='CRIADA', data=None, criada='2025-01-02',
                 al_os='ALTESTE', desc='antiga', ativos=['X (chave no mapa)'], metodo='entre ativos', km_desc=None, m=1.0)
    some = dict(velha, os='000002/2025', os_completa='ETO-RD-PA 000002/2025')
    lp = dict(regra_anos={'T1': 3, 'T2': 4, 'T3': 5}, fonte='velha.xlsx', gerado='2026-01-01', equipes=[],
              os=[velha, some], vaos={ids[0]: [0], ids[1]: [0, 1]}, sem_local=[])
    (site / 'alimentadores' / 'ALTESTE' / 'limpeza.js').write_text('window.LIMPEZA = ' + json.dumps(lp) + ';\n', encoding='utf-8')
    cab = ['NUMERO_OS', 'SITUACAO', 'DESCRICAO_OS', 'ALIMENTADOR', 'DATA_CRIACAO', 'INICIO_EXEC', 'FIM_EXEC', 'ESTRUTURA',
           'ELEMENTO', 'ABRANGENCIA', 'MONT_ABRANGENCIA']
    h = lambda *a: dt.datetime(*a)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Dados'
    for r in [cab, ['ETO-RD-PA 000001/2025', 'EXECUTADA', 'nova descrição', 'ALTESTE', h(2025, 1, 2), h(2026, 2, 1), h(2026, 2, 1)] + [None] * 4,
              ['ETO-RD-PA 000003/2026', 'EXECUTADA', 'limpeza a partir da chave', 'ALTESTE', h(2026, 1, 5), h(2026, 3, 1), h(2026, 3, 1),
               None, '7900001', None, None],
              ['ETO-LD-PCM 000004/2026', 'EXECUTADA', 'chave', 'ALTESTE', h(2026, 1, 5), h(2026, 3, 1), h(2026, 3, 1), None, '7900001', None, None]]:
        ws.append(r)
    (d / 'bases' / 'limpeza').mkdir(parents=True)
    wb.save(d / 'bases' / 'limpeza' / 'OS teste.xlsx')
    limpeza_faixa.usa_site(site)
    limpeza_faixa.main(str(d / 'bases' / 'limpeza' / 'OS teste.xlsx'))
    out = inspecao.le_js(site / 'alimentadores' / 'ALTESTE' / 'limpeza.js')
    por = {o['os_completa']: k for k, o in enumerate(out['os'])}
    assert set(por) == {'ETO-RD-PA 000001/2025', 'ETO-RD-PA 000003/2026'}, por
    o1 = out['os'][por['ETO-RD-PA 000001/2025']]
    assert o1['sit'] == 'EXECUTADA' and o1['data'] == '2026-02-01' and o1['desc'] == 'nova descrição', o1
    assert o1['metodo'] == 'entre ativos' and o1['ativos'] == ['X (chave no mapa)'], 'posição publicada mantida'
    assert sorted(v for v, ks in out['vaos'].items() if por['ETO-RD-PA 000001/2025'] in ks) == sorted(ids), 'mesmos vãos'
    o3 = out['os'][por['ETO-RD-PA 000003/2026']]
    assert o3['ativos'] == ['7900001 (chave no mapa)'] and o3['metodo'] == 'resto do trecho a partir do ativo', o3
    print('ok limpeza_reaproveita')


def bases_pelo_menu(d):
    """atualiza_bases: tipo pela pasta, Crítica sem mês recusada, inspeção processada e dados/bases.json."""
    assert atualiza_bases.tipo_de('bases/inspecao/BASE INSPEÇÃO.xlsx') == 'inspecao'
    assert atualiza_bases.tipo_de('outra/inspecao/x.xlsx') is None and atualiza_bases.tipo_de('bases/LEIA-ME.md') is None
    import zipfile
    with tempfile.TemporaryDirectory() as tmp:
        z = Path(tmp) / 'crit.zip'
        with zipfile.ZipFile(z, 'w') as zz:
            zz.writestr('pasta/../Critica-CHEIO_09-2026.txt', 'x')
            zz.writestr('outro.txt', 'y')
        txt = atualiza_bases.arquivos_critica([z], tmp)
        assert [t.name for t in txt] == ['Critica-CHEIO_09-2026.txt'] and txt[0].parent == Path(tmp), 'zip: só Critica*.txt, sem caminho'
        sem = Path(tmp) / 'Critica-sem-mes.txt'
        sem.write_text('x')
        try:
            atualiza_bases.arquivos_critica([sem], tmp)
            assert False, 'Crítica sem MM-AAAA no nome tem que parar'
        except SystemExit as e:
            assert 'MM-AAAA' in str(e)
    (d / 'bases' / 'inspecao').mkdir(parents=True)
    shutil.copyfile(d / 'insp.xlsx', d / 'bases' / 'inspecao' / 'insp.xlsx')
    feitos = atualiza_bases.atualiza(d / 'site', [d / 'bases' / 'inspecao' / 'insp.xlsx', d / 'bases' / 'inspecao' / 'LEIA-ME.md'],
                                     hoje='2026-10-13')
    assert list(feitos) == ['inspecao']
    est = json.loads((d / 'site' / 'dados' / 'bases.json').read_text(encoding='utf-8'))
    assert est['inspecao']['arquivo'] == 'insp.xlsx' and est['inspecao']['enviado'] == '2026-10-13' and est['inspecao']['os'] == 3, est
    assert est['limpeza']['os'] == 2 and est['critica']['nome'] == 'Crítica (NAE)', est
    assert not list((d / 'site').rglob('__pycache__')), 'nada de __pycache__ no site'
    print('ok bases_pelo_menu')


def manutencao_no_site(d):
    """OS de manutenção da inspeção: canceladas fora; poste pela ESTRUTURA; sem poste no mapa fica sem posição; pela
    automação (bases/manutencao) entra no bases.json."""
    import datetime as dt
    import openpyxl
    cab = ['INSP_NUMERO_OS', 'INSP_FIM_EXEC', 'NUMERO_OS', 'DESCRICAO_OS', 'ALIMENTADOR', 'CRITICIDADE', 'CODEQUIPE',
           'DATA_CRIACAO', 'FIM_EXEC', 'DESCESQUEMA', 'SITUACAO', 'ESTRUTURA', 'DATA_LIMITE', 'EQUIPE_EXECUTORA']
    h = lambda *a: dt.datetime(*a)
    linhas = [
        ['ETO-INPA01  000002/2026', h(2026, 1, 1), 'ETO-RD-PA  000010/2026', 'CRUZETA > RACHADA', 'ALTESTE', 'ALTA', 'ETO-RD-PA',
         h(2026, 2, 1), None, 'MC - CRUZETA', 'CRIADA', '190-N4', h(2026, 3, 1), None],
        ['ETO-INPA01  000002/2026', h(2026, 1, 1), 'ETO-RD-PA  000011/2026', 'ISOLADOR > QUEBRADO', 'ALTESTE', 'BAIXA', 'ETO-RD-PA',
         h(2026, 2, 1), h(2026, 2, 20), 'MC - ISOLADOR', 'EXECUTADA', '190-N4', h(2026, 8, 1), 'DEOP'],
        ['ETO-INPA01  000002/2026', h(2026, 1, 1), 'ETO-RD-PA  000012/2026', 'POSTE > INCLINADO', 'ALTESTE', 'BAIXA', 'ETO-RD-PA',
         h(2026, 2, 2), None, 'MC - POSTES', 'CRIADA', '190-999999', h(2027, 1, 1), None],
        ['ETO-INPA01  000002/2026', h(2026, 1, 1), 'ETO-RD-PA  000013/2026', 'CRUZETA > RACHADA', 'ALTESTE', 'BAIXA', 'ETO-RD-PA',
         h(2026, 2, 3), None, 'MC - CRUZETA', 'CANCELADA', '190-N4', h(2026, 3, 1), None],
        ['ETO-INPA01  000005/2026', h(2026, 1, 1), 'ETO-RD-PA  000014/2026', 'CRUZETA', 'ALOUTRO', 'BAIXA', 'ETO-RD-PA',
         h(2026, 2, 4), None, 'MC - CRUZETA', 'CRIADA', '190-1', h(2026, 3, 1), None],
    ]
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in [cab] + linhas:
        ws.append(r)
    (d / 'bases' / 'manutencao').mkdir(parents=True)
    xlsx = d / 'bases' / 'manutencao' / 'OS geradas.xlsx'
    wb.save(xlsx)
    feitos = atualiza_bases.atualiza(d / 'site', [xlsx], hoje='2026-10-13')
    assert list(feitos) == ['manutencao']
    m = inspecao.le_js(d / 'site' / 'alimentadores' / 'ALTESTE' / 'manutencao.js')
    por = {o['os']: o for o in m['os']}
    assert set(por) == {'000010/2026', '000011/2026', '000012/2026'}, 'cancelada e alimentador sem página ficam de fora'
    dados = inspecao.le_js(d / 'site' / 'alimentadores' / 'ALTESTE' / 'dados.js')
    n4 = next([v[0], v[1]] if v[6] == 'N4' else [v[2], v[3]] for v in dados['vaos'] if 'N4' in (v[6], v[7]))
    o = por['000010/2026']
    assert o['ll'] == n4 and o['poste'] == 'N4' and o['tipo'] == 'CRUZETA' and o['limite'] == '2026-03-01', o
    assert o['insp'] == 'ETO-INPA01 000002/2026' and o['exec'] is None, o
    assert por['000011/2026']['exec'] == '2026-02-20' and por['000011/2026']['executora'] == 'DEOP'
    assert por['000012/2026']['ll'] is None, 'poste que não está no mapa: sem posição'
    est = json.loads((d / 'site' / 'dados' / 'bases.json').read_text(encoding='utf-8'))
    assert est['manutencao']['os'] == 3 and est['manutencao']['arquivo'] == 'OS geradas.xlsx', est['manutencao']
    pag = (d / 'site' / 'alimentadores' / 'ALTESTE' / 'index.html').read_text(encoding='utf-8')
    assert '<script src="manutencao.js"></script>' in pag and (d / 'site' / 'assets' / 'manutencao.js').exists()
    print('ok manutencao_no_site')


def executadas_no_site(d):
    """OS executadas (card Manutenção): só MC/MP; posição pelo poste da ESTRUTURA, pelo ativo da ESTRUTURA ou pelo
    ELEMENTO (índice de ativos do site); ativo de outro alimentador vai para a página dele; sem posição fica na lista;
    a OS repetida conta uma vez; pela automação (bases/executadas, .zip) entra no bases.json."""
    import zipfile
    site = d / 'site'
    # um poste com número no índice (os postes da rede de teste são N1, N2..., que ficam fora do índice)
    fp = site / 'dados' / 'ativos' / (gera_trechos.fatia('426869') + '.json')
    idx = json.loads(fp.read_text(encoding='utf-8')) if fp.exists() else {}
    idx['426869'] = [['ALTESTE', 'poste', 'T2', 'T2', -7.1985, -48.2]]
    fp.write_text(json.dumps(idx), encoding='utf-8')
    cab = ['NUMERO_OS', 'DESCRICAO_OS', 'ALIMENTADOR', 'CRITICIDADE', 'CODEQUIPE', 'FIM_EXEC', 'DESCESQUEMA', 'SITUACAO',
           'ESTRUTURA', 'ELEMENTO']
    linhas = [
        ['ETO-RD-PA  000001/2025', 'TROCA DE\n\nISOLADOR   N1', 'ALTESTE', 'ALTA', 'ETO-RD-PA', '10/03/2025 14:00:00', 'MC - ISOLADOR',
         'EXECUTADA', '190-426869', '8800001'],
        ['ETO-RD-PA  000002/2026', 'TROCA DE TRAFO', 'ALTESTE', '', 'ETO-RD-PA', '05/02/2026 09:00:00',
         'MC - SUBSTITUIÇÃO DE TRANSFORMADOR', 'EXECUTADA', 'GRU-57-5700001', '5700001'],
        ['ETO-RD-PA  000003/2024', 'PODA', 'ALTESTE', '', 'DBM-PD', '01/07/2024 08:00:00', 'MP - PODA DE ARVORE', 'EXECUTADA',
         '190-999999', '07900001'],
        ['DBM-PDGU02  000004/2024', 'INSPEÇÃO', 'ALTESTE', '', 'DBM-PD', '01/07/2024 08:00:00', 'INSP - INSPEÇÃO PODA', 'EXECUTADA',
         '190-426869', ''],
        ['ETO-RD-PA  000005/2023', 'CRUZETA', 'ALOUTRO', '', 'ETO-RD-PA', '02/02/2023 10:00:00', 'MC - CRUZETA', 'EXECUTADA',
         '190-426869', ''],
        ['ETO-RD-PA  000006/2023', 'POSTE', 'ALTESTE', '', 'ETO-RD-PA', '03/03/2023 10:00:00', 'MC - POSTES', 'EXECUTADA',
         '190-111', ''],
        ['ETO-RD-PA  000007/2023', 'CRUZETA', 'ALSEMPAG', '', 'ETO-RD-PA', '04/03/2023 10:00:00', 'MC - CRUZETA', 'EXECUTADA',
         '190-222', ''],
        ['ETO-RD-PA  000001/2025', 'TROCA DE\n\nISOLADOR   N1', 'ALTESTE', 'ALTA', 'ETO-RD-PA', '10/03/2025 14:00:00', 'MC - ISOLADOR',
         'EXECUTADA', '190-426869', '8800001'],
        ['ETO-RD-PA  000008/2023', 'X', 'ALTESTE', '', 'ETO-RD-PA', '04/03/2023 10:00:00', 'MC - CRUZETA', 'CANCELADA', '190-426869', ''],
    ]
    import csv
    import io
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=';', lineterminator='\r\n')
    for r in [cab] + linhas:
        w.writerow(r)
    (d / 'bases' / 'executadas').mkdir(parents=True)
    z = d / 'bases' / 'executadas' / 'OS_STATUS_Executadas_teste.zip'
    with zipfile.ZipFile(z, 'w', zipfile.ZIP_DEFLATED) as zz:
        zz.writestr('OS_STATUS_Executadas_teste.txt', '﻿' + buf.getvalue())
    feitos = atualiza_bases.atualiza(site, [z], hoje='2026-10-14')
    assert list(feitos) == ['executadas']
    ex = inspecao.le_js(site / 'alimentadores' / 'ALTESTE' / 'executadas.js')
    por = {o[0]: o for o in ex['os']}
    assert list(por) == ['ETO-RD-PA 000002/2026', 'ETO-RD-PA 000001/2025', 'ETO-RD-PA 000003/2024', 'ETO-RD-PA 000006/2023',
                         'ETO-RD-PA 000005/2023'], 'só MC/MP executadas, uma vez cada, da mais recente para a mais antiga'
    at = lambda o: ex['ativos'][o[4]] if o[4] >= 0 else None
    o1 = por['ETO-RD-PA 000001/2025']
    assert at(o1) == ['426869', 'poste', 'T2', -7.1985, -48.2, 'p'] and o1[2] == '2025-03-10' and o1[5] == 'TROCA DE ISOLADOR N1', o1
    assert ex['tipos'][o1[1]] == 'MC - ISOLADOR' and o1[6] == 'alta' and o1[7] == ''
    assert at(por['ETO-RD-PA 000002/2026'])[:2] == ['5700001', 'TR'] and at(por['ETO-RD-PA 000002/2026'])[5] == 'e', 'trafo da ESTRUTURA'
    assert at(por['ETO-RD-PA 000003/2024'])[:2] == ['07900001', '79'] and at(por['ETO-RD-PA 000003/2024'])[5] == 'l', \
        'poste fora do mapa: o ELEMENTO, com o código como está na OS'
    o5 = por['ETO-RD-PA 000005/2023']
    assert o5[4] == o1[4] and o5[7] == 'ALOUTRO', 'OS de alimentador sem página, ativo aqui: o mesmo ponto, alimentador anotado'
    assert por['ETO-RD-PA 000006/2023'][4] == -1, 'sem posição: fica na lista'
    assert len(ex['ativos']) == 3 and (ex['de'], ex['ate']) == ('2023-02-02', '2026-02-05'), ex
    est = json.loads((site / 'dados' / 'bases.json').read_text(encoding='utf-8'))
    assert est['executadas']['os'] == 5 and est['executadas']['ate'] == '2026-02-05', est['executadas']
    assert est['executadas']['arquivo'] == 'OS_STATUS_Executadas_teste.zip' and est['executadas']['nome'] == 'OS de manutenção executadas'
    pag = (site / 'alimentadores' / 'ALTESTE' / 'index.html').read_text(encoding='utf-8')
    assert '<script src="executadas.js"></script>' in pag and (site / 'assets' / 'executadas.js').exists()
    def busca_os(os_):                              # índice da busca de OS: posição do ativo da OS, não do tipo
        k = gera_trechos.chave_os(os_)
        fx = json.loads((site / 'dados' / 'os' / f'{gera_trechos.fatia_os(k)}.json').read_text(encoding='utf-8'))
        return next(e for e in fx[k] if e[2] == 'e')
    assert busca_os('ETO-RD-PA 000001/2025')[4:] == [-7.1985, -48.2], busca_os('ETO-RD-PA 000001/2025')
    assert busca_os('ETO-RD-PA 000006/2023')[4:] == [None, None], 'sem posição na busca também'
    print('ok executadas_no_site')


def inspecao_no_site(d):
    """Base de inspeção: só EXECUTADA; caminho do ELEMENTO subindo até a ABRANGÊNCIA; só o ponto; sem posição."""
    import datetime as dt
    import openpyxl
    cab = ['NUMERO_OS', 'ALIMENTADOR', 'SITUACAO', 'CODEQUIPE', 'CRITICIDADE', 'INICIO_EXEC', 'FIM_EXEC',
           'ELEMENTO', 'ABRANGENCIA', 'MONT_ABRANGENCIA', 'DESCRICAO_OS']
    h = lambda *a: dt.datetime(*a)
    linhas = [
        ['ETO-INPA01  000001/2026', 'ALTESTE', 'EXECUTADA', 'ETO-INPA01', 'ALTA', h(2026, 1, 10, 8), h(2026, 1, 10, 10),
         '0300001', '7900001', 'DJ_ALTESTE_FICT', 'OS despachada automaticamente pelo Scheduler/SOMA'],
        ['ETO-INPA01  000002/2026', 'ALTESTE', 'EXECUTADA', 'ETO-INPA01', 'BAIXA', h(2026, 1, 1, 8), h(2026, 4, 1, 8),
         '7900001', '9000001', 'DJ_ALTESTE_FICT', 'BLOCO CHAVE FACA - 9000001'],
        ['ETO-INPA01  000003/2026', 'ALTESTE', 'EXECUTADA', 'ETO-INPA01', 'BAIXA', h(2026, 2, 1, 8), h(2026, 2, 1, 9),
         '9000002', '9000003', '', ''],
        ['ETO-INPA01  000004/2026', 'ALTESTE', 'CRIADA', 'ETO-INPA01', 'BAIXA', None, None, '0300001', '7900001', '', ''],
        ['ETO-INPA01  000005/2026', 'ALOUTRO', 'EXECUTADA', 'ETO-INPA01', 'BAIXA', h(2026, 3, 1, 8), h(2026, 3, 1, 9),
         '1', '2', '', ''],
    ]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Dados'
    for r in [cab] + linhas:
        ws.append(r)
    wb.save(d / 'insp.xlsx')
    r = inspecao.main(d / 'insp.xlsx', d / 'site')
    assert r['executadas'] == 4 and r['sem_pagina'] == {'ALOUTRO': 1}, 'só EXECUTADA; alimentador fora do site à parte'
    assert r['situacao'] == {'no mapa (bloco)': 2, 'sem posição': 1}, r['situacao']
    arq = d / 'site' / 'alimentadores' / 'ALTESTE' / 'inspecao.js'
    ins = inspecao.le_js(arq)
    os_ = {o['os']: o for o in ins['os']}
    c = json.loads((d / 'site' / 'dados' / 'alimentadores' / 'ALTESTE.json').read_text(encoding='utf-8'))
    poste = {ch['nome']: ch['poste'] for ch in c['chaves']}
    assert (poste['0300001'], poste['7900001']) == ('N6', 'N3'), poste
    vaos_de = lambda o: sorted(v for v, ks in ins['vaos'].items() if ins['os'].index(o) in ks)
    dados = inspecao.le_js(d / 'site' / 'alimentadores' / 'ALTESTE' / 'dados.js')
    abaixo_n6 = sorted(v[5] for v in dados['vaos'] if v[6] in ('N6', 'N7') and v[7] in ('N6', 'N7'))
    o1 = os_['000001/2026']                           # ELEMENTO: o fusível (N6); bloco abaixo dele, sem chave depois
    assert o1['metodo'].startswith('bloco abaixo do ELEMENTO') and vaos_de(o1) == abaixo_n6 and o1['fins'] == [], (o1, abaixo_n6)
    assert o1['data'] == '2026-01-10' and o1['desc'] == '' and o1['pa'] and o1['pe'], o1
    o2 = os_['000002/2026']                           # ELEMENTO: o religador (N3); bloco até o fusível, que fecha o bloco
    assert vaos_de(o2) == ['TR-004', 'TR-005', 'TR-006'] and o2['m'] == 100.2, (vaos_de(o2), o2)
    assert [f[2] for f in o2['fins']] == ['0300001'] and o2['pa'] is None, o2
    assert o2['data'] == '2026-01-01', 'FIM mais de 60 dias depois do INICIO: vale o INICIO'
    o3 = os_['000003/2026']
    assert o3['metodo'].startswith('sem posição') and o3['pt'] is None and ins['base_ate'] == '2026-03-01'
    antes = arq.read_bytes()
    inspecao.main(d / 'insp.xlsx', d / 'site')
    assert arq.read_bytes() == antes, 'inspeção: segunda execução igual'
    print('ok inspecao_no_site')


CRITICA_CAB = ('DTA_CMPT_HDE;TIPO_INTER_KTI;COD_ELE_REDE_PROBLEMA;COD_ELE_PROBLEMA;ABRANGÊNCIA;NUM_SEQ_OPER_INIC_HDE;'
               'DTA_ABERT;NUM_PASSO_INIC_HDE;COD_ELE_REDE_INTERROMPIDO;COD_ELE_INTERROMPIDO;NUM_PASSO_FIM_HDE;DTA_FECH;'
               'COD_ELE_FECHADO;QTD_CONS_INTER_FAT;DURACAO;DES_CAUSA_INTER_CAU;DES_SUB_CAUSA_INTER_SCR;QTD_IFEC_ALIM;'
               'QTD_IDEC_ALIM;COD_ALIMENTADOR')


def critica_sem_kml(d):
    """Crítica nova só com os dados do site: chave de início, chave fora da topologia, trafo do GIS e ativo desconhecido."""
    site = d / 'site'
    ativos_gis.grava_trafos(site, ativos_gis.transformadores(d / 'gis')[0])
    linha = lambda n, abr, tipo, prob, inter, dia: ';'.join((
        '01/01/26', 'COLETIVA' if tipo != 'UC' else 'INDIVIDUAL', tipo, prob, abr, str(n), f'{dia}/01/2026 10:00', '1',
        'TR' if tipo == 'UC' else tipo, inter, '2', f'{dia}/01/2026 11:00', inter, '3', '60', 'MEIO AMBIENTE',
        'ARVORE DE GRANDE PORTE NA REDE', '0', '0', 'ALTESTE'))
    arq = d / 'Critica-CHEIO_01-2026.txt'
    arq.write_text('\n'.join([CRITICA_CAB, linha(1, 'DJ', 'DJ', 'DJ_ALTESTE_FICT', 'DJ_ALTESTE_FICT', '02'),
                              linha(2, 'RL', 'RL', '7900001', '7900001', '03'),
                              linha(3, 'CH FUS', 'CH FUS', '3100001', '3100001', '04'),
                              linha(4, 'TR', 'UC', '123456', '5700001', '05'),
                              linha(5, 'CH FUS', 'CH FUS', '9999999', '9999999', '06')]) + '\n', encoding='latin-1')
    xlsx = d / 'NAE.xlsx'
    assert atualiza_critica.main(['--site', str(site), '--critica', str(arq), '--planilha', str(xlsx)]) == 0 and xlsx.exists()
    c = json.loads((site / 'dados' / 'alimentadores' / 'ALTESTE.json').read_text(encoding='utf-8'))
    got = [(o['ativo'], o['trecho'], o['regra'], o['fora']) for o in c['nae']['oc']]
    assert got == [('DJ_ALTESTE_FICT', 'T1', 'abrangência DJ', None),
                   ('7900001', 'T2', 'chave de início do trecho', None),
                   ('3100001', 'T2', 'vão mais próximo do ativo', None),
                   ('5700001', 'T2', ativos_gis.REGRA, None),
                   ('9999999', None, 'ativo fora do KML', 'fora do cadastro atual')], got
    assert c['nae']['meses'] == ['2026-01'] and c['origem']['critica'] == ['Critica-CHEIO_01-2026.txt']
    assert ativos_gis.le_js(site / 'alimentadores' / 'ALTESTE' / 'dados.js')[1]['nae'] == c['nae'], 'página igual ao JSON'
    assert atualiza_critica.main(['--site', str(site), '--conferir']) == 0, 'refazer a associação dá o mesmo'
    antes = assinatura(site)
    assert atualiza_critica.main(['--site', str(site), '--critica', str(arq)]) == 0
    assert assinatura(site) == antes, 'rodar a mesma Crítica de novo não muda nada'


def ativos_gis_no_site(d):
    """Trafo do GIS (camada Transformadores) dá trecho à NAE que estava 'ativo fora do KML'."""
    site, y = d / 'site', lambda i: -7.2 + i * 0.0003
    (d / 'gis0').mkdir()
    assert ativos_gis.aplica(site, d / 'gis0')[0]['nae_localizadas'] == 0, 'site sem Crítica (nae vazio) não quebra'
    assert ativos_gis.outros(site, pts={})['nae_hoje_em_outro'] == 0 and ativos_gis.posicoes(site, pts={})
    oc = dict(mes='2026-01', dia='02/01 10:00', col=True, abr='TR', ativo='5700001', prob='5700001', causa='Meio ambiente',
              sub='Árvore', prog=False, cons=3, dur=60, trecho=None, fora='fora do cadastro atual', regra='ativo fora do KML',
              dist_m=None)
    jf, mf = site / 'dados' / 'alimentadores' / 'ALTESTE.json', site / 'alimentadores' / 'ALTESTE' / 'dados.js'
    c = json.loads(jf.read_text(encoding='utf-8'))
    c['nae'] = dict(meses=['2026-01'], oc=[dict(oc), dict(oc, ativo='5799999', prob='5799999')])
    jf.write_text(json.dumps(c), encoding='utf-8')
    cab, D = ativos_gis.le_js(mf)
    D['nae'] = json.loads(json.dumps(c['nae']))
    mf.write_text(cab + json.dumps(D) + ';\n', encoding='utf-8')
    pm = lambda n, lat, lon: (f'<Placemark><name>{n}</name><ExtendedData><Data name="Alimentador"><value>ALTESTE</value></Data>'
                              f'</ExtendedData><Point><coordinates>{lon},{lat},0</coordinates></Point></Placemark>')
    (d / 'gis').mkdir()
    (d / 'gis' / 'Transformadores.kml').write_text('<kml><Document><name>Transformadores</name>' + pm('5700001', y(4.5), -48.20001)
                                                   + pm('5700002', y(4.5), -48.21) + '</Document></kml>', encoding='utf-8')
    tot, _ = ativos_gis.aplica(site, d / 'gis')
    assert tot['nae_localizadas'] == 1 and tot['nae_sem_ponto'] == 1, tot
    c = json.loads(jf.read_text(encoding='utf-8'))
    o = c['nae']['oc'][0]
    assert (o['trecho'], o['regra']) == ('T2', ativos_gis.REGRA) and o['dist_m'] < 2, 'trafo a 1 m do vão t4 (T2)'
    assert c['nae']['oc'][1]['trecho'] is None, 'trafo que não está no GIS continua sem trecho'
    assert [a[4] for a in c['ativos_kml']] == ['T2', None], 'trafo a ~1 km da rede não entra em trecho'
    assert ativos_gis.le_js(mf)[1]['nae']['oc'][0]['trecho'] == 'T2', 'página (dados.js) igual ao JSON'
    assert gera_trechos.main(['--saida', str(site), '--so-interface']) == 0
    antes = (jf.read_text(encoding='utf-8'), mf.read_text(encoding='utf-8'))
    ativos_gis.aplica(site, d / 'gis'); ativos_gis.outros(site, d / 'gis'); ativos_gis.posicoes(site, d / 'gis')
    c2 = json.loads(jf.read_text(encoding='utf-8'))
    assert c2['nae'] == json.loads(antes[0])['nae'] | {'oc': c2['nae']['oc']} and c2['nae']['oc'][0]['trecho'] == 'T2', 'rodar de novo não muda'
    assert c2['nae']['oc'][0].get('lonlat') and 'hoje' not in c2['nae']['oc'][0]
    hit = json.loads((site / 'dados' / 'ativos' / (gera_trechos.fatia('5700001') + '.json')).read_text(encoding='utf-8'))['5700001']
    assert hit[0][:3] == ['ALTESTE', 'TR', 'T2'], 'busca de ativo acha o trafo no trecho'


def test_chave_sem_segmento():
    """GIS com a linha passando direto pela chave (sem o conector): a chave entra no vão onde está o ponto dela."""
    import gis_para_daimon as g
    segs = {((0.0, 0.0), (100.0, 0.0)): dict(nome='v1', fase='ABC')}
    ll, grau = {(0.0, 0.0): (-48.0, -10.0), (100.0, 0.0): (-47.9991, -10.0)}, {(0.0, 0.0): 1, (100.0, 0.0): 1}
    k = g.encaixa((50.3, 0.2), segs, ll, grau, '0300001')
    assert k == ((49.8, 0.0), (50.8, 0.0)) and segs[k]['nome'] == 'Suporte Chave 0300001' and len(segs) == 3
    assert abs(sum(math.dist(a, b) for a, b in segs) - 100) < 1e-6, 'a rede não muda de tamanho'
    assert g.encaixa((50.0, 5.0), {((0.0, 0.0), (100.0, 0.0)): dict(nome='v', fase='')}, dict(ll), dict(grau), 'X') is None


def test_equipes_limpeza():
    assert limpeza_faixa.conta('EXECUTADA', 'ETO-RD-GU 004972/2022'), 'executada da equipe do polo conta'
    assert not limpeza_faixa.conta('EXECUTADA', 'ETO-LD-PCM 000452/2025'), 'executada de outra equipe não conta'
    assert limpeza_faixa.conta('CRIADA', 'ETO-LD-PCM 000452/2025'), 'pendente aparece, de qualquer equipe'



def test_chaves_js():
    c = dict(chaves=[dict(nome='DJ X', tipo='DJ', lonlat=[-48.1, -10.1]), dict(nome='7900000001', tipo='79', lonlat=[-48.2, -10.2]),
                     dict(nome='0300000001', tipo='03', lonlat=[-48.3, -10.3]), dict(nome='8800000001', tipo='88', lonlat=[-48.4, -10.4]),
                     dict(nome='8000000001', tipo='80', lonlat=[-48.5, -10.5]), dict(nome='0300000002', tipo='SEC', lonlat=[-48.6, -10.6])])
    js = gera_trechos.chaves_js(c)
    assert js.startswith('window.CHAVES = ') and js.endswith(';\n')
    cs = json.loads(js[len('window.CHAVES = '):-2])
    assert cs == [['8800000001', '88', -10.4, -48.4], ['8000000001', '80', -10.5, -48.5]], 'só as que não dividem, em [lat, lon]'



def test_redes_mapa():
    """Rede da prévia do município: vãos encadeados numa linha por galho, nas coordenadas do mapa (PROJ_MUN)."""
    a, b, c_, d = gera_trechos.PROJ_MUN
    pts = [(-48.30, -10.20), (-48.31, -10.20), (-48.32, -10.20), (-48.31, -10.21)]      # tronco com uma derivação no meio
    vaos = [[*pts[0], *pts[1]], [*pts[1], *pts[2]], [*pts[1], *pts[3]]]
    with tempfile.TemporaryDirectory() as dd:
        saida = Path(dd)
        gera_trechos.redes_mapa(saida, [dict(alimentador='AL1', validacao=dict(status='ok'), vaos=vaos),
                                         dict(alimentador='AL2', validacao=dict(status='falha'), vaos=vaos)])
        r = json.loads((saida / 'dados' / 'redes' / 'AL1.json').read_text(encoding='utf-8'))
        assert not (saida / 'dados' / 'redes' / 'AL2.json').exists(), 'alimentador com falha fica de fora'
        assert r['d'].count('M') == 3, 'três galhos a partir da derivação: ' + r['d']
        x = lambda lon: round(a * lon + b, 2)
        y = lambda lat: round(c_ * lat + d, 2)
        assert r['b'] == [x(-48.32), y(-10.20), x(-48.30), y(-10.21)], r['b']


def test_chave_os():
    """Número/ano da OS: com ou sem equipe e zeros à esquerda; a fatia bate com a conta do busca-ativo.js."""
    for t in ('ETO-RD-PO 002507/2026', '2507/2026', '002507 / 2026'):
        assert gera_trechos.chave_os(t) == '2507/2026', t
    assert gera_trechos.chave_os('sem número') is None and gera_trechos.chave_os('2507/26') is None
    assert gera_trechos.fatia_os('2507/2026') == f'{sum(ord(c) * 31 ** i for i, c in enumerate(reversed("2507/2026"))) % 2 ** 32 % 64:02x}'


def test_manutencao_polos():
    cfg = json.loads((gera_trechos.PKG / 'site' / 'polos.json').read_text(encoding='utf-8'))
    pa, pb = cfg['polos'][0], cfg['polos'][1]
    with tempfile.TemporaryDirectory() as d:
        saida = Path(d)
        (saida / 'assets').mkdir()
        ok = dict(status='ok')
        dados = [dict(alimentador='AL1', validacao=ok, trechos=[dict(municipios={pa['municipios'][0]: 900.0, pb['municipios'][0]: 100.0})]),
                 dict(alimentador='AL2', validacao=ok, trechos=[dict(municipios={pa['municipios'][1]: 50.0})]),
                 dict(alimentador='AL3', validacao=dict(status='falha'), trechos=[dict(municipios={pa['municipios'][0]: 10.0})])]
        for al, oss in [('AL1', [('EXECUTADA', '2026-01-01'), ('CRIADA', '2026-01-10'), ('CRIADA', '2026-01-10')]),
                        ('AL2', [('PROGRAMADA', '2099-01-01')])]:
            (saida / 'alimentadores' / al).mkdir(parents=True)
            (saida / 'alimentadores' / al / 'manutencao.js').write_text('window.MANUTENCAO = ' + json.dumps(
                dict(base_ate='2026-10-10', os=[dict(sit=a, limite=b) for a, b in oss])) + ';\n', encoding='utf-8')
        obj = gera_trechos.manutencao_polos(saida, dados)
        assert obj['al'] == {'AL1': pa['id'], 'AL2': pa['id']}, 'cada alimentador no polo de mais rede; falha fica de fora'
        q = obj['polos'][pa['id']]
        assert (q['als'], q['exec'], q['lim']) == (2, 1, {'2026-01-10': 2, '2099-01-01': 1}), q
        assert obj['polos'][pb['id']]['als'] == 0
        assert (saida / 'assets' / 'manutencao-polos.js').read_text(encoding='utf-8').startswith('window.MANUT_POLOS = {')


if __name__ == '__main__':
    test_pipeline()
    print('ok test_pipeline')
    test_equipes_limpeza()
    print('ok test_equipes_limpeza')
    test_chave_sem_segmento()
    print('ok test_chave_sem_segmento')
    test_chaves_js()
    print('ok test_chaves_js')
    test_redes_mapa()
    print('ok test_redes_mapa')
    test_chave_os()
    print('ok test_chave_os')
    test_manutencao_polos()
    print('ok test_manutencao_polos')
