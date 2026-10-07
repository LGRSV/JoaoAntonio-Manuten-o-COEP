"""Testes das regras e das saídas.  python -m alimentador.testes   (ou: pytest alimentador/testes.py)"""
import contextlib
import io
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from . import critica, exemplo, exporta, regras, trechos


def rede_linha(tipos, desvio=None, desvio2=None):
    """Linha reta DJ → tipos[0] → tipos[1] ... com vãos de ~40 m. desvio/desvio2 = (índice, [tipos]) criam ramais
    (nós d1, d2… para um lado e e1, e2… para o outro)."""
    nos = {'DJ': dict(tipo='DJ', nome='DJ', poste='P0', lat=-9.6, lon=-48.45)}
    vaos, ant = [], 'DJ'
    for i, t in enumerate(tipos, 1):
        nos[f'n{i}'] = dict(tipo=t, nome=f'{t}-{i}' if t != 'P' else f'P{i}', poste=f'P{i}', lat=-9.6, lon=-48.45 + i * 0.00037)
        vaos.append((ant, f'n{i}'))
        ant = f'n{i}'
    for letra, dv, lado in (('d', desvio, -1), ('e', desvio2, 1)):
        if not dv:
            continue
        idx, ts = dv
        ant = f'n{idx}'
        for j, t in enumerate(ts, 1):
            k = f'{letra}{j}'
            nos[k] = dict(tipo=t, nome=f'{t}-{k}' if t != 'P' else f'P{k}', poste=f'P{k}', lat=-9.6 + lado * j * 0.00037,
                          lon=-48.45 + idx * 0.00037)
            vaos.append((ant, k))
            ant = k
    return nos, vaos


def zona_apos(res, nome):
    return next(a['zona_depois'] for a in res['ativos'] if a['nome'] == nome)


def test_zonas():
    res = regras.analisar(*rede_linha(['P', 'P', 'P', 'P'], desvio=(1, ['03', 'P'])), 'DJ')
    assert zona_apos(res, '03-d1') == 3, 'fusível em derivação do T1 vai direto para T3'
    res = regras.analisar(*rede_linha(['79', 'P', '79', 'P', 'P', 'P', 'P', 'P', 'P'], desvio=(4, ['03', 'P', '03', 'P'])), 'DJ')
    assert [zona_apos(res, n) for n in ('79-1', '79-3', '03-d1', '03-d3')] == [2, 2, 3, 3], 'religador nunca cria T3; não existe T4'
    res = regras.analisar(*rede_linha(['79', 'P', '03', 'P', '79', 'P']), 'DJ')
    assert zona_apos(res, '03-3') == 2 and zona_apos(res, '79-5') == 2, '03 antes de religador é seccionamento'
    assert res['fim_tronco'] == '79-5' and any('seccionamento' in a for a in res['avisos'])
    assert not any(v['tronco'] and v['zona'] > 2 for v in res['vaos'])
    res = regras.analisar(*rede_linha(['P', '79', 'P', '03', 'P', '03', 'P']), 'DJ')
    assert [zona_apos(res, n) for n in ('03-4', '03-6')] == [2, 3], 'só a 1ª chave fusível depois do religador é do tronco'
    assert [a['tipo'] for a in res['ativos'] if a['nome'] in ('03-4', '03-6')] == ['FTR', '03']
    assert any('depois do religador do tronco' in a for a in res['avisos'])
    res = regras.analisar(*rede_linha(['P', '79', 'P', '79', 'P'], desvio=(1, ['P', '79', 'P', '03', 'P'])), 'DJ')
    assert zona_apos(res, '79-d2') == 2, 'religador em ramal também conta'
    assert zona_apos(res, '03-d4') == 3, 'fusível em ramal fora da linha principal é T3'
    res = regras.analisar(*rede_linha(['79', '88', 'P', '79', 'P']), 'DJ')
    assert zona_apos(res, '88-2') == 2, 'chave 88 não muda a zona'


def test_tronco_e_blocos():
    res = regras.analisar(*rede_linha(['P', '79', 'P', '79', 'P', 'P'], desvio=(3, ['P', '03', 'P'])), 'DJ')
    tronco = [v for v in res['vaos'] if v['tronco']]
    assert len(tronco) == 4 and all(v['zona'] <= 2 for v in tronco), 'tronco vai do DJ ao 2º religador'
    assert res['fim_tronco'] == '79-4'
    ramais = {v['ramal'] for v in res['vaos'] if not v['tronco']}
    assert ramais == {'Ramal após 79-4', 'Ramal P3'}, ramais
    assert not res['avisos']


def test_tronco_sem_segundo_religador():
    res = regras.analisar(*rede_linha(['P', '79', 'P', 'P', 'P'], desvio=(3, ['03', 'P', 'P', 'P', 'P', 'P', 'P'])), 'DJ')
    assert any('sem 2º religador' in a for a in res['avisos'])
    assert next(a['tipo'] for a in res['ativos'] if a['nome'] == '03-d1') == 'FTR', '1ª chave fusível depois do religador'
    assert all(v['zona'] <= 2 for v in res['vaos']) and all(v['tronco'] for v in res['vaos'] if v['_c'].startswith('d'))
    assert not any(v['tronco'] for v in res['vaos'] if v['_c'] in ('n4', 'n5')), 'ponta mais curta fica fora do tronco'
    res = regras.analisar(*rede_linha(['03', 'P', 'P']), 'DJ')
    assert not any(v['tronco'] for v in res['vaos']) and {v['zona'] for v in res['vaos']} == {1, 3}, 'DJ direto num 03: sem religador antes, abre T3'


def test_fusivel_no_tronco_300m():
    """1º fusível depois do religador na linha principal: a menos de 300 m dele não abre T3; a 300 m ou mais, abre."""
    perto = regras.analisar(*rede_linha(['P', '79', 'P', 'P', '03', 'P', 'P']), 'DJ')            # 03 a ~120 m do 79
    assert next(a['tipo'] for a in perto['ativos'] if a['nome'] == '03-5') == 'FTR' and zona_apos(perto, '03-5') == 2
    longe = regras.analisar(*rede_linha(['P', '79'] + ['P'] * 8 + ['03', 'P', 'P']), 'DJ')         # 03 a ~365 m
    assert next(a['tipo'] for a in longe['ativos'] if a['nome'] == '03-11') == '03' and zona_apos(longe, '03-11') == 3
    assert not any('1ª(s) chave(s) fusível(is)' in a for a in longe['avisos'])


def test_rede_invalida():
    nos, vaos = rede_linha(['P', 'P'])
    for ruim, msg in ((vaos + [('DJ', 'n2')], 'dois vãos'),
                      (vaos, None)):
        if msg is None:
            nos2 = dict(nos, solto=dict(tipo='P', nome='X', poste='X', lat=0, lon=0))
            try:
                regras.analisar(nos2, ruim, 'DJ')
                raise AssertionError('nó solto deveria falhar')
            except regras.RedeInvalida as e:
                assert 'sem ligação' in str(e)
        else:
            try:
                regras.analisar(nos, ruim, 'DJ')
                raise AssertionError('deveria falhar')
            except regras.RedeInvalida as e:
                assert msg in str(e)


def test_exemplo_completo():
    nos, vaos = exemplo.montar_rede()
    res = regras.analisar(nos, vaos, 'DJ')
    assert len(res['blocos']) == 15, len(res['blocos'])
    assert abs(sum(v['ext'] for v in res['vaos']) - res['total_m']) < 1e-6
    assert abs(sum(b['ext'] for b in res['blocos']) - res['total_m']) < 1e-6
    for v in res['vaos']:
        assert v['caminho'][-1] == v['ref'] and v['dist_ref'] >= 0
    for b in res['blocos']:
        assert {v['zona'] for v in res['vaos'] if v['bloco'] == b['n']} == {b['zona']}
    assert not any(b['ref'].startswith('88') for b in res['blocos']), 'chave 88 não é ativo de referência'
    b88 = {c: b for b in res['blocos'] for c in b['chaves88']}
    assert b88['88-A']['ref'] == 'RL1' and b88['88-A']['tronco']
    assert b88['88-B']['ref'] == '03-E'
    zonas_ref = {b['ref']: b['zona'] for b in res['blocos']}
    assert zonas_ref['RL2'] == 2 and zonas_ref['RL3'] == 2, 'depois do 2º religador ainda é T2 (T2 Ramal)'
    assert zonas_ref['03-E'] == 3 and zonas_ref['03-F'] == 3 and zonas_ref['03-C'] == 3
    rl2 = next(b for b in res['blocos'] if b['ref'] == 'RL2' and b['tronco'])
    assert set(rl2['ate']) == {'RL3', '03-E'}, 'o T depois do 2º religador é tronco até os próximos ativos'
    ramais = {b['ref']: b['ramal'] for b in res['blocos'] if not b['tronco']}
    assert ramais['RL3'] == 'Ramal após RL3' and ramais['03-E'] == 'Ramal após 03-E' and ramais['03-F'] == 'Ramal após 03-E'
    assert not any(v['tronco'] and v['zona'] > 2 for v in res['vaos'])

    with tempfile.TemporaryDirectory() as d:
        d = Path(d)
        exporta.kmz(res, 'T', d / 't.kmz')
        exporta.excel(res, 'T', d / 't.xlsx')
        exporta.mapa(res, 'T', d / 't.html')
        root = ET.fromstring(zipfile.ZipFile(d / 't.kmz').read('doc.kml'))
        ns = {'k': 'http://www.opengis.net/kml/2.2'}
        linhas = [p for p in root.iter('{http://www.opengis.net/kml/2.2}Placemark') if p.find('.//k:LineString', ns) is not None]
        assert len(linhas) == len(res['vaos'])
        wb = load_workbook(d / 't.xlsx')
        assert wb.sheetnames == ['Leia-me', 'Árvore de blocos', 'Consulta', 'Blocos', 'Por bloco', 'Trechos', 'Ativos']
        assert wb['Trechos'].max_row == len(res['vaos']) + 1
        html = (d / 't.html').read_text(encoding='utf-8')
        assert '%%' not in html and '"trechos"' in html


CAB_CRITICA = ['DTA_CMPT_HDE', 'TIPO_INTER_KTI', 'COD_ELE_REDE_PROBLEMA', 'COD_ELE_PROBLEMA', 'ABRANGÊNCIA',
               'NUM_SEQ_OPER_INIC_HDE', 'DTA_ABERT', 'COD_ELE_REDE_INTERROMPIDO', 'COD_ELE_INTERROMPIDO',
               'QTD_CONS_INTER_FAT', 'DURACAO', 'DES_CAUSA_INTER_CAU', 'DES_SUB_CAUSA_INTER_SCR', 'QTD_IFEC_ALIM',
               'QTD_IDEC_ALIM', 'COD_ALIMENTADOR']


def _kml_daimon(chaves, trechos):
    pm_c = ''.join(f'<Placemark><name>{n}</name><styleUrl>#mapDaimonChave{"NA" if na else ""}</styleUrl>'
                   f'<Point><coordinates>{lon},{lat},0</coordinates></Point></Placemark>' for n, lat, lon, na in chaves)
    pm_t = ''.join(f'<Placemark><name>{n}</name><LineString><coordinates>{a[1]},{a[0]},0 {b[1]},{b[0]},0'
                   f'</coordinates></LineString></Placemark>' for n, a, b in trechos)
    return ('<?xml version="1.0" encoding="UTF-8"?><kml xmlns="http://www.opengis.net/kml/2.2"><Document>'
            f'<Folder><name>Circ_X</name><Folder><name>Chave</name>{pm_c}</Folder>'
            f'<Folder><name>Trecho</name>{pm_t}</Folder></Folder></Document></kml>')


def test_leitor_daimon():
    """DJ → religador com 90 em série e bypass NA em paralelo → fusível. A energia passa pelo religador."""
    from . import daimon
    y = lambda i: -9.6 + i * 0.0003     # pontos numa linha, ~33 m entre eles
    P = {i: (y(i), -48.45) for i in range(12)}
    B = (y(4), -48.4497)                # ponto do desvio (bypass) ao lado
    chaves = [('DJ_X_FICT', (y(-1) + P[0][0]) / 2, (-48.4503 + P[0][1]) / 2, False),   # DJ no meio do suporte (empate)
              ('9000001', (y(2) + y(3)) / 2, -48.45, False),
              ('7900001', (y(3) + y(4)) / 2, -48.45, False), ('9000002', (y(4) + y(5)) / 2, -48.45, False),
              ('8000001', (y(3) + B[0]) / 2, -48.4498, True), ('0300001', (y(8) + y(9)) / 2, -48.45, False)]
    trechos = [('Suporte Chave DJ_X_FICT', (y(-1), -48.4503), P[0]), ('t1', P[0], P[1]), ('t2', P[1], P[2]),
               ('Suporte Chave 9000001', P[2], P[3]), ('Suporte Chave 7900001', P[3], P[4]),
               ('Suporte Chave 9000002', P[4], P[5]), ('Suporte Chave 8000001', P[3], B), ('t3', B, P[5]),
               ('t4', P[5], P[6]), ('t5', P[6], P[7]), ('t6', P[7], P[8]), ('Suporte Chave 0300001', P[8], P[9]),
               ('t7', P[9], P[10]), ('t8', P[10], P[11])]
    with tempfile.TemporaryDirectory() as d:
        arq = Path(d) / 'SE-X_Circ-ALTESTE.kml'
        arq.write_text(_kml_daimon(chaves, trechos), encoding='utf-8')
        nos, vaos, raiz, info = daimon.ler_circuito(arq)
    assert info['abertas'] == ['8000001'], info['abertas']
    assert nos[raiz]['lat'] == round(P[0][0], 6), 'raiz é a ponta do suporte do DJ mais longe do ponto do DJ'
    tipos = sorted(n['tipo'] for n in nos.values() if n['tipo'] != 'P')
    assert tipos == ['03', '79', 'DJ'], tipos          # 90 vira passagem; 80 NA fica fora
    res = regras.analisar(nos, vaos, raiz)
    assert zona_apos(res, '7900001') == 2 and zona_apos(res, '0300001') == 2, 'fusível na linha principal é do tronco'
    # sai só o suporte do DJ e o da chave NA; o fio do desvio fica como ponta sem saída, ligado do lado de baixo
    assert len(res['vaos']) == len(trechos) - 2
    assert info['nomes_vaos'][next(v for v in vaos if nos[v[1]]['tipo'] == '79')] == 'Suporte Chave 7900001'



def rede_trechos():
    """DJ → RL1 → RL2 → tronco até P11; entre RL1 e RL2 sai um desvio com religador e fusível; depois do RL2 sai um
    ramal com fusível dentro de fusível."""
    return rede_linha(['P', '79', 'P', '79', 'P', 'P', 'P', 'P', 'P', 'P', 'P'], desvio=(3, ['P', '79', 'P', '03', 'P']),
                      desvio2=(5, ['03', 'P', '03', 'P']))


def test_trechos():
    nos, vaos = rede_trechos()
    res = regras.analisar(nos, vaos, 'DJ')
    ts = {t['nome']: t for t in trechos.segmentar(nos, res)}
    assert sorted(ts) == ['T1', 'T2', 'T2-A', 'T2-B', 'T3-A', 'T3-B', 'T3-C'], sorted(ts)
    assert (ts['T1']['inicio'], ts['T1']['fim']) == ('DJ', ['79-2'])
    assert (ts['T2']['inicio'], ts['T2']['fim']) == ('79-2', ['79-4']), 'T2 vai de religador até religador'
    assert (ts['T2-A']['inicio'], ts['T2-A']['pai'], ts['T2-A']['tronco']) == ('79-4', 'T2', True), '2º religador abre T2-A'
    assert ts['T2']['derivacoes'] == ['T2-B'], 'religador fora do tronco é derivação'
    assert (ts['T2-B']['inicio'], ts['T2-B']['pai'], ts['T2-B']['derivacoes']) == ('79-d2', 'T2', ['T3-B'])
    assert ts['T2-A']['fim'] == ['fim de linha · poste P11'] and ts['T2-B']['fim'] == ['03-d4'], 'sem religador: linha principal'
    assert ts['T2-A']['derivacoes'] == ['T3-A'] and ts['T3-B']['fim'] == ['fim de linha · poste Pd5']
    assert [ts[n]['inicio'] for n in ('T3-A', 'T3-B', 'T3-C')] == ['03-e1', '03-d4', '03-e3'], 'só T3, letras pela distância do DJ'
    assert ts['T3-A']['fim'] == ['03-e3'] and ts['T3-C']['pai'] == 'T3-A', 'fusível dentro de fusível também é T3'
    assert abs(sum(t['ext'] for t in ts.values()) - res['total_m']) < 0.5
    assert all(v['zona'] == ts[v['trecho']]['zona'] for v in res['vaos'])
    assert [trechos.letras(i) for i in (0, 25, 26, 27, 51, 52)] == ['A', 'Z', 'AA', 'AB', 'AZ', 'BA']
    with tempfile.TemporaryDirectory() as d:
        trechos.mapa(nos, res, 'TESTE', Path(d) / 'm.html')
        html = (Path(d) / 'm.html').read_text(encoding='utf-8')
        assert '%%' not in html and '"nome":"T2-A"' in html
        trechos.kmz(nos, res, 'TESTE', Path(d) / 't.kmz')
        doc = ET.fromstring(zipfile.ZipFile(Path(d) / 't.kmz').read('doc.kml'))
        n = '{http://www.opengis.net/kml/2.2}'
        pastas = [f.find(n + 'name').text for g in doc.find(n + 'Document').find(n + 'Folder').findall(n + 'Folder')
                  for f in g.findall(n + 'Folder')]
        assert [p.split(' · ')[0] for p in pastas] == ['T1', 'T2', 'T2-A', 'T2-B', 'T3-A', 'T3-B', 'T3-C'], pastas
        assert pastas[4].startswith('T3-A · 03-e1 → 03-e3')



def test_gerar_trechos():
    """Comando de ponta a ponta: KML do Daimon + Crítica → mapa, KMZ, página do site e índice por SE."""
    from . import gerar_trechos
    y = lambda i: -9.6 + i * 0.0003
    P = {i: (y(i), -48.45) for i in range(8)}
    meio = lambda i: ((y(i) + y(i + 1)) / 2, -48.45)
    chaves = [('DJ_ALTESTE_FICT', *meio(-1), False), ('7900001', *meio(2), False), ('0300001', *meio(5), False)]
    trechos_kml = [('Suporte Chave DJ_ALTESTE_FICT', (y(-1), -48.45), P[0]), ('t1', P[0], P[1]), ('t2', P[1], P[2]),
                   ('Suporte Chave 7900001', P[2], P[3]), ('t3', P[3], P[4]), ('t4', P[4], P[5]),
                   ('Suporte Chave 0300001', P[5], P[6]), ('t5', P[6], P[7])]
    linhas = [';'.join(CAB_CRITICA),
              ';'.join(['01/01/26', 'COLETIVO', 'DJ', '1R0', 'DJ', '1', '11/01/2026 18:16', 'DJ', '1R0', '300', '51', 'X', 'Y', ',1', ',1', 'ALTESTE']),
              ';'.join(['01/01/26', 'COLETIVO', 'CH', '0300001', 'CH FUS', '2', '20/01/2026 10:00', 'TR', 'TRX', '40', '90', 'X', 'Y', ',1', ',1', 'ALTESTE']),
              ';'.join(['01/01/26', 'COLETIVO', 'DJ', '1R0', 'SE', '3', '22/01/2026 10:00', 'DJ', '1R0', '300', '2', 'X', 'Y', ',1', ',1', 'ALTESTE']),
              ';'.join(['01/01/26', 'INDIVIDUAL', 'UC', '999', 'UC', '4', '23/01/2026 10:00', 'TR', 'TR-SUMIU', '1', '9', 'X', 'Y', ',1', ',1', 'ALTESTE'])]
    with tempfile.TemporaryDirectory() as d:
        d = Path(d)
        kml = d / 'SE-XTST_Circ-ALTESTE.kml'
        kml.write_text(_kml_daimon(chaves, trechos_kml), encoding='utf-8')
        (d / 'Critica-1.txt').write_bytes(('\r\n'.join(linhas) + '\r\n').encode('latin-1'))
        (d / 'site').mkdir()
        saida = io.StringIO()
        with contextlib.redirect_stdout(saida):
            gerar_trechos.main([str(kml), '--saida', str(d / 'saida'), '--critica', str(d / 'Critica-*.txt'),
                                '--mes', '2026-01', '--site', str(d / 'site'), '--se-nome', 'SE-XTST=Teste'])
        texto = saida.getvalue()
        assert 'T1: DJ ALTESTE → 7900001' in texto and '    T1: 1 · 0' in texto and '    T2: 1 · 0' in texto, texto
        assert '    fora · na subestação: 1 · 0' in texto and '    fora do cadastro atual: 0 · 1' in texto, texto
        assert sorted(p.name for p in (d / 'saida' / 'ALTESTE').iterdir()) == ['ALTESTE-trechos.kmz', 'mapa-trechos.html']
        pag = (d / 'site' / 'ALTESTE.html').read_text(encoding='utf-8')
        assert 'noindex' in pag and 'href="ALTESTE-trechos.kmz"' in pag and '← Todos os alimentadores' in pag
        assert 'cdnjs' not in pag and 'Leaflet 1.9.4' in pag, 'site com o Leaflet embutido'
        indice = (d / 'site' / 'index.html').read_text(encoding='utf-8')
        assert 'SE-XTST · Teste' in indice and 'href="ALTESTE.html"' in indice and '3</strong> coletivas' in indice
        assert (d / 'site' / 'ALTESTE-trechos.kmz').exists() and (d / 'site' / '.nojekyll').exists()


def test_copia_da_skill():
    """A skill leva uma cópia do código (scripts/alimentador) para rodar fora do repositório: tem que ser igual."""
    raiz = Path(__file__).resolve().parent
    copia = raiz.parent / '.claude' / 'skills' / 'trechos-alimentador' / 'scripts' / 'alimentador'
    arquivos = ['__init__.py', 'daimon.py', 'regras.py', 'trechos.py', 'critica.py', 'gerar_trechos.py',
                'mapa_trechos.html', 'site_indice.html', 'leaflet/leaflet.js', 'leaflet/leaflet.css', 'leaflet/LICENSE']
    diferentes = [a for a in arquivos if (raiz / a).read_bytes() != (copia / a).read_bytes()]
    assert not diferentes, f'copie de novo para {copia}: {diferentes}'


def test_critica():
    nos, vaos = rede_trechos()
    res = regras.analisar(nos, vaos, 'DJ')
    ts = trechos.segmentar(nos, res)
    cab = CAB_CRITICA
    lin = lambda *v: ';'.join(v)
    dj = lin('01/01/26', 'COLETIVO', 'DJ', '9DJ', 'DJ', '1', '11/01/2026 18:16', 'DJ', '9DJ', '300', '51', 'TERCEIROS', 'OBJETO', ',07', ',06', 'ALX')
    arq1 = [';'.join(cab), dj,
            lin('01/01/26', 'COLETIVO', 'DJ', '9DJ', 'DJ', '1', '11/01/2026 18:16', 'DJ', '9DJ', '20', '2', 'TERCEIROS', 'OBJETO', ',01', ',00', 'ALX'),
            lin('01/01/26', 'COLETIVO', 'CH', '03-e1', 'CH FUS', '2', '20/01/2026 10:00', 'TR', 'TRX', '40', '90', 'NAO-PROGRAMADAS MEIO AMBIENTE', 'ARVORE', ',01', ',02', 'ALX'),
            lin('01/01/26', 'INDIVIDUAL', 'UC', '123456', 'UC', '3', '21/01/2026 09:00', 'TR', 'TR-B', '1', '60', 'CONEXAO', 'RAMAL', ',0002', ',0001', 'ALX'),
            lin('01/02/26', 'INDIVIDUAL', 'UC', '654321', 'UC', '4', '02/02/2026 09:00', 'TR', 'TR-SUMIU', '1', '30', 'CONEXAO', 'RAMAL', ',0002', ',0001', 'ALX'),
            lin('01/01/26', 'COLETIVO', 'DJ', '9DJ', 'DJ', '5', '12/01/2026 08:00', 'DJ', '9DJ', '99', '5', 'X', 'Y', ',1', ',1', 'OUTRO'),
            lin('01/01/26', 'COLETIVO', 'DJ', '9DJ', 'SE', '6', '25/01/2026 07:00', 'DJ', '9DJ', '300', '2', 'DISTRIBUICAO PROGRAMADA', 'Z', ',1', ',1', 'ALX')]
    # no arquivo de fevereiro: a linha de janeiro repetida e o passo de fevereiro de novo, com o indicador recalculado
    arq2 = [';'.join(cab), dj,
            lin('01/02/26', 'INDIVIDUAL', 'UC', '654321', 'UC', '4', '02/02/2026 09:00', 'TR', 'TR-SUMIU', '1', '30', 'CONEXAO', 'RAMAL', ',0005', ',0001', 'ALX')]
    fim_b = next(v for v in res['vaos'] if v['trecho'] == 'T2-B')['b']
    with tempfile.TemporaryDirectory() as d:
        arqs = [Path(d, 'Critica-CHEIO_01-2026.txt'), Path(d, 'Critica-CHEIO_02-2026.txt')]
        for arq, linhas in zip(arqs, (arq1, arq2)):
            arq.write_bytes(('\r\n'.join(linhas) + '\r\n').encode('latin-1'))
        linhas = critica.ler(arqs, 'ALX')
    assert len(linhas) == 6, 'só o alimentador pedido e cada passo uma vez'
    oc = critica.por_trecho(critica.ocorrencias(linhas), res, ts, {'TR-B': fim_b}, {'ALY': {'TR-SUMIU': (0, 0)}})
    dj_oc = next(o for o in oc if o['abr'] == 'DJ')
    assert (dj_oc['trecho'], dj_oc['cons'], dj_oc['mes']) == ('T1', 320, '2026-01'), 'passos somados numa ocorrência'
    fus = next(o for o in oc if o['abr'] == 'CH FUS')
    assert fus['trecho'] == 'T3-A' and not fus['programada'], 'fusível que abre trecho; NAO-PROGRAMADAS não é programada'
    assert not critica.arvore_na_rede(fus), 'NAO-PROGRAMADAS MEIO AMBIENTE não é a causa Meio Ambiente'
    assert critica.arvore_na_rede(dict(causa='Meio ambiente', sub='Arvore de grande porte na rede / na faixa'))
    assert critica.arvore_na_rede(dict(causa='Meio ambiente', sub='Eucalipto na rede / na faixa'))
    assert not critica.arvore_na_rede(dict(causa='Entrada de servico aereo', sub='Arvore no ramal'))
    se = next(o for o in oc if o['abr'] == 'SE')
    assert (se['trecho'], se['fora'], se['programada']) == (None, 'na subestação', True), 'subestação fica à parte'
    uc = [o for o in oc if o['abr'] == 'UC']
    assert uc[0]['trecho'] == 'T2-B' and uc[0]['prob'] == '' and not uc[0]['coletiva'], 'individual pelo trafo, sem o número da UC'
    assert (uc[1]['mes'], uc[1]['trecho'], uc[1]['fora']) == ('2026-02', None, 'hoje no ALY'), 'competência e ativo fora do cadastro'
    assert (uc[1]['cons'], uc[1]['fec']) == (1, .0005), 'passo em dois arquivos conta uma vez, com o do mês da competência'


if __name__ == '__main__':
    testes = [f for n, f in sorted(globals().items()) if n.startswith('test_')]
    for f in testes:
        f()
        print('ok', f.__name__)
    print(f'{len(testes)} testes passaram')
