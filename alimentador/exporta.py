"""Saídas a partir do resultado de regras.analisar(): KMZ, Excel e mapa HTML."""
import json
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.table import Table, TableStyleInfo

TIPO = {'DJ': 'Disjuntor (DJ)', '79': 'Religador (79)', '03': 'Chave fusível (03)', '88': 'Chave 88'}
REGRAS_TXT = [
    ('Zona', 'T1 do DJ ao 1º religador (79). Depois de qualquer religador é T2, inclusive religador em ramal: '
             'religador nunca cria T3. Depois de uma chave fusível (03) é no mínimo T3; se já estava em T3 ou mais, soma uma.'),
    ('Tronco', 'Caminho DJ → 1º religador → 2º religador; termina no 2º religador. Nunca tem T3. '
               'O resto é ramal, com o nome do poste onde sai do tronco.'),
    ('Ativo de referência', 'Último DJ, 79 ou 03 antes do vão, vindo do DJ. Ele engloba o vão e tudo o que vem depois.'),
    ('Ativos posteriores englobados', 'Os DJ/79/03 que ficam depois do vão ou do bloco. Se o ativo de referência abrir, todos ficam sem energia.'),
    ('Bloco', 'Vãos com o mesmo ativo de referência e o mesmo ramal (de um ativo até o próximo). '
              'Chaves 88 não dividem bloco: aparecem listadas na coluna "Chaves 88 no bloco".'),
    ('Nível 1…N', 'Sequência de ativos desde o DJ até o ativo de referência. Filtrar um nível mostra tudo o que fica depois daquele ativo.'),
    ('Distâncias', 'Calculadas a partir das coordenadas (haversine, raio 6.371 km), seguindo a rede.'),
]
ZFILL = ['DBEAFE', 'DCFCE7', 'FEF3C7', 'FFEDD5', 'F3E8FF', 'FCE7F3', 'E0E7FF', 'CCFBF1']
KML_Z = ['fff6823b', 'ff5ec522', 'ff15ccfa', 'ff1673f9', 'ffef46d9', 'ff9948ec', 'fff755a8', 'ffa6d414']  # aabbggrr
zfill = lambda z: PatternFill('solid', fgColor=ZFILL[min(z, len(ZFILL)) - 1])


def _prox(res):
    return {b['n']: b['ate'] for b in res['blocos']}


# ============================================================ KMZ
def kmz(res, nome, caminho, defeito=None, aviso_exemplo=False):
    icon = 'http://maps.google.com/mapfiles/kml/paddle/wht-circle.png'
    dot = 'http://maps.google.com/mapfiles/kml/shapes/placemark_circle.png'
    cor_ativo = {'DJ': 'ff0000ff', '79': 'ff00ff00', '03': 'ff00ffff', '88': 'ffff0000'}
    prox = _prox(res)
    desc = ('<b>DADOS DE EXEMPLO.</b> Circuito ilustrativo; não é rede real.<br>' if aviso_exemplo else '') + \
        '<br>'.join(f'<b>{a}:</b> {escape(b)}' for a, b in REGRAS_TXT[:5])
    k = ['<?xml version="1.0" encoding="UTF-8"?>', '<kml xmlns="http://www.opengis.net/kml/2.2"><Document>',
         f'<name>{escape(nome)} — zonas, ramais e blocos</name><open>1</open>',
         f'<description><![CDATA[{desc}]]></description>']
    for z in range(1, len(KML_Z) + 1):
        for t, w in (('Tronco', 6), ('Ramal', 3)):
            k.append(f'<Style id="z{z}{t}"><LineStyle><color>{KML_Z[z - 1]}</color><width>{w}</width></LineStyle></Style>')
    for t, c in cor_ativo.items():
        k.append(f'<Style id="at{t}"><IconStyle><color>{c}</color><scale>{0.8 if t == "88" else 1.1}</scale>'
                 f'<Icon><href>{icon}</href></Icon></IconStyle><LabelStyle><scale>0.85</scale></LabelStyle></Style>')
    k.append(f'<Style id="poste"><IconStyle><color>ffcccccc</color><scale>0.45</scale><Icon><href>{dot}</href></Icon></IconStyle>'
             '<LabelStyle><scale>0</scale></LabelStyle></Style>')
    k.append('<Style id="defeito"><IconStyle><color>ff2e10c8</color><scale>1.3</scale>'
             '<Icon><href>http://maps.google.com/mapfiles/kml/shapes/caution.png</href></Icon></IconStyle></Style>')

    def xd(campos):
        return '<ExtendedData>' + ''.join(
            f'<Data name="{escape(a)}"><value>{escape(str(b))}</value></Data>' for a, b in campos) + '</ExtendedData>'

    def coord(p):
        return f'{p[1]:.7f},{p[0]:.7f},0'

    def placemark_vao(v):
        z = min(v['zona'], len(KML_Z))
        campos = [('Trecho', v['id']), ('Poste início', v['ini']), ('Poste fim', v['fim']),
                  ('Extensão (m)', f"{v['ext']:.1f}"), ('Distância do DJ (m)', f"{v['dist_dj']:.0f}"),
                  ('Distância do ativo de referência (m)', f"{v['dist_ref']:.0f}"),
                  ('Zona', f"T{v['zona']}"), ('Tronco/Ramal', 'Tronco' if v['tronco'] else 'Ramal'),
                  ('Ramal', v['ramal']), ('Bloco', v['bloco']), ('Ativo de referência', v['ref']),
                  ('Tipo do ativo de referência', TIPO[v['ref_tipo']]), ('Próximo(s) ativo(s)', ' / '.join(prox[v['bloco']])),
                  ('Ativos posteriores englobados', ', '.join(v['posteriores']) or '—'),
                  ('Chave 88 no fim do vão', v['chave88_fim'] or '—'), ('Caminho desde o DJ', ' › '.join(v['caminho']))]
        d = f"Ref.: {v['ref']} · T{v['zona']} · {v['ramal']} · " + ' › '.join(v['caminho'])
        return (f'<Placemark><name>{v["id"]}</name><description>{escape(d)}</description>'
                f'<styleUrl>#z{z}{"Tronco" if v["tronco"] else "Ramal"}</styleUrl>{xd(campos)}'
                f'<LineString><tessellate>1</tessellate><coordinates>{coord(v["a"])} {coord(v["b"])}</coordinates>'
                '</LineString></Placemark>')

    def pasta_bloco(b):
        vs = [v for v in res['vaos'] if v['bloco'] == b['n']]
        c88 = f" · 88: {', '.join(b['chaves88'])}" if b['chaves88'] else ''
        out = [f'<Folder><name>Bloco {b["n"]} · T{b["zona"]} · {escape(b["de"])} → {escape(" / ".join(b["ate"]))} '
               f'({len(vs)} vãos · {b["ext"]:.0f} m{escape(c88)})</name>']
        out += [placemark_vao(v) for v in vs]
        out.append('</Folder>')
        return out

    k.append(f'<Folder><name>{escape(nome)}</name><open>1</open>')
    k.append('<Folder><name>Tronco (DJ → 2º religador)</name><open>1</open>')
    for b in res['blocos']:
        if b['tronco']:
            k += pasta_bloco(b)
    k.append('</Folder><Folder><name>Ramais</name><open>1</open>')
    for ram in dict.fromkeys(b['ramal'] for b in res['blocos'] if not b['tronco']):
        k.append(f'<Folder><name>{escape(ram)}</name>')
        for b in res['blocos']:
            if b['ramal'] == ram:
                k += pasta_bloco(b)
        k.append('</Folder>')
    k.append('</Folder>')
    k.append('<Folder><name>Ativos (DJ, religadores, chaves)</name>')
    for a in res['ativos']:
        campos = [('Ativo', a['nome']), ('Tipo', TIPO[a['tipo']]), ('Poste', a['poste']),
                  ('Zona antes', f"T{a['zona_antes']}" if a['zona_antes'] else '—'), ('Zona depois', f"T{a['zona_depois']}"),
                  ('Bloco onde está', a['bloco'] or '—'),
                  ('Blocos em que é referência', ', '.join(map(str, a['blocos_proprios'])) or '—'),
                  ('Ativos posteriores englobados', ', '.join(a['posteriores']) or '—')]
        k.append(f'<Placemark><name>{escape(a["nome"])}</name><styleUrl>#at{a["tipo"]}</styleUrl>{xd(campos)}'
                 f'<Point><coordinates>{a["lon"]:.7f},{a["lat"]:.7f},0</coordinates></Point></Placemark>')
    k.append('</Folder><Folder><name>Postes</name>')
    vistos = set()
    for v in res['vaos']:
        for pid, p in ((v['ini'], v['a']), (v['fim'], v['b'])):
            if pid not in vistos:
                vistos.add(pid)
                k.append(f'<Placemark><name>{escape(pid)}</name><styleUrl>#poste</styleUrl>'
                         f'<Point><coordinates>{coord(p)}</coordinates></Point></Placemark>')
    k.append('</Folder>')
    if defeito:
        v = next(x for x in res['vaos'] if x['id'] == defeito)
        m = ((v['a'][0] + v['b'][0]) / 2, (v['a'][1] + v['b'][1]) / 2)
        d = (f"{v['id']} entre {v['ini']} e {v['fim']} · Ref.: {v['ref']} · T{v['zona']} · {v['ramal']} · "
             + ' › '.join(v['caminho']) + f" · {v['dist_ref']:.0f} m depois do {v['ref']}")
        k.append(f'<Folder><name>Exemplo de defeito</name><Placemark><name>Defeito (exemplo) · {v["id"]}</name>'
                 f'<description>{escape(d)}</description><styleUrl>#defeito</styleUrl>'
                 f'<Point><coordinates>{coord(m)}</coordinates></Point></Placemark></Folder>')
    k.append('</Folder></Document></kml>')
    with zipfile.ZipFile(caminho, 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('doc.kml', '\n'.join(k))


# ============================================================ EXCEL
def excel(res, nome, caminho, defeito=None, aviso_exemplo=False):
    F = 'Arial'
    body = Font(name=F, size=10)
    bold = Font(name=F, size=10, bold=True)
    hdr_font, hdr_fill = Font(name=F, bold=True, color='FFFFFF'), PatternFill('solid', fgColor='2B3136')
    thin = Side(style='thin', color='D5D9D6')
    link = Font(name=F, size=10, color='1F4E9E', underline='single')
    vaos, blocos, ativos = res['vaos'], res['blocos'], res['ativos']
    prox = _prox(res)
    NV = max(len(b['caminho']) for b in blocos)           # níveis de ativos no caminho
    MAXN = max(len(b['ate']) for b in blocos) or 1
    MAXP = max(len(b['posteriores']) for b in blocos) or 1
    NT = len(vaos)
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True

    def cabecalho(ws, headers, linha=1):
        for j, h in enumerate(headers, 1):
            c = ws.cell(linha, j, h)
            c.font, c.fill, c.alignment = hdr_font, hdr_fill, Alignment(vertical='center', wrap_text=True)
        ws.row_dimensions[linha].height = 32

    def tabela(ws, nome_tab, headers, rows, widths):
        cabecalho(ws, headers)
        for r in rows:
            ws.append(r)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = body
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[L(i)].width = w
        t = Table(displayName=nome_tab, ref=f'A1:{L(len(headers))}{len(rows) + 1}')
        t.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True)
        ws.add_table(t)

    # --- Leia-me
    ws = wb.active
    ws.title = 'Leia-me'
    ws['A1'] = f'Zonas, ramais e blocos — {nome}'
    ws['A1'].font = Font(name=F, bold=True, size=14)
    linhas = ([('DADOS DE EXEMPLO', 'Circuito ilustrativo. Não é rede real. Serve para validar o formato.')] if aviso_exemplo else []) \
        + REGRAS_TXT + [('Abas', 'Árvore de blocos: visão rápida da sequência, agrupada por tronco e ramal. '
                                 'Consulta: digite um poste e veja onde ele está. Blocos: um bloco por linha, com os ativos '
                                 'posteriores um por coluna e juntos numa coluna só. Por bloco: os vãos agrupados embaixo de '
                                 'cada bloco (botão + à esquerda). Trechos: um vão por linha. Ativos: resumo por ativo.')] \
        + [('Aviso', a) for a in res['avisos']]
    for i, (a, b) in enumerate(linhas, 3):
        ws.cell(i, 1, a).font = bold
        ws.cell(i, 2, b).font = body
        ws.cell(i, 1).alignment = Alignment(vertical='top')
        ws.cell(i, 2).alignment = Alignment(wrap_text=True, vertical='top')
    if aviso_exemplo:
        ws['A3'].font = Font(name=F, bold=True, size=10, color='C8102E')
    ws.column_dimensions['A'].width, ws.column_dimensions['B'].width = 28, 100
    r0 = len(linhas) + 4
    ws.cell(r0, 1, 'Cores das zonas').font = bold
    for z in range(1, max(v['zona'] for v in vaos) + 1):
        c = ws.cell(r0 + z, 1, f'T{z}')
        c.fill, c.font = zfill(z), body

    # --- Trechos
    wt = wb.create_sheet('Trechos')
    HT = ['Seq', 'Trecho', 'Poste início', 'Poste fim', 'Extensão (m)', 'Distância do DJ até o fim (m)',
          'Distância do ativo de referência até o fim (m)', 'Zona', 'Tronco/Ramal', 'Ramal', 'Bloco',
          'Ativo de referência', 'Tipo do ativo de referência', 'Próximo(s) ativo(s)', 'Ativos posteriores englobados',
          'Qtd. ativos posteriores', 'Chave 88 no fim do vão'] + [f'Nível {i}' for i in range(1, NV + 1)] + \
         ['Caminho desde o DJ', 'Latitude (meio do vão)', 'Longitude (meio do vão)', 'Mapa', 'Chave do caminho (uso das fórmulas)']
    col = {h: i + 1 for i, h in enumerate(HT)}
    rows = []
    for i, v in enumerate(vaos, 1):
        m = ((v['a'][0] + v['b'][0]) / 2, (v['a'][1] + v['b'][1]) / 2)
        v['_meio'] = m
        rows.append([i, v['id'], v['ini'], v['fim'], v['ext'], v['dist_dj'], v['dist_ref'], f"T{v['zona']}",
                     'Tronco' if v['tronco'] else 'Ramal', v['ramal'], v['bloco'], v['ref'], TIPO[v['ref_tipo']],
                     ' / '.join(prox[v['bloco']]), ', '.join(v['posteriores']) or '—', len(v['posteriores']),
                     v['chave88_fim'] or None, *(v['caminho'] + [None] * (NV - len(v['caminho']))),
                     ' › '.join(v['caminho']), round(m[0], 7), round(m[1], 7), 'abrir', '|' + '|'.join(v['caminho']) + '|'])
    tabela(wt, 'Trechos', HT, rows, [6, 9, 11, 11, 10, 13, 15, 7, 11, 17, 7, 12, 18, 22, 34, 11, 12] + [8] * NV
           + [34, 13, 13, 8, 26])
    wt.freeze_panes = 'C2'
    for i, v in enumerate(vaos, 2):
        wt.cell(i, col['Zona']).fill = zfill(v['zona'])
        c = wt.cell(i, col['Mapa'])
        c.hyperlink, c.font = f"https://www.google.com/maps?q={v['_meio'][0]:.7f},{v['_meio'][1]:.7f}", link
        wt.cell(i, col['Extensão (m)']).number_format = '0.0'
        for h in ('Distância do DJ até o fim (m)', 'Distância do ativo de referência até o fim (m)'):
            wt.cell(i, col[h]).number_format = '#,##0'
        if v['id'] == defeito:
            for j in range(1, len(HT) + 1):
                wt.cell(i, j).fill = PatternFill('solid', fgColor='FFE4E6')
    wt.cell(1, col['Distância do ativo de referência até o fim (m)']).comment = Comment(
        'Quanto andar a partir do ativo de referência, seguindo a rede, até o fim deste vão.', 'Claude')
    rng = lambda h: f"Trechos!${L(col[h])}$2:${L(col[h])}${NT + 1}"
    R_BLOCO, R_EXT, R_REF, R_POSTE, R_CHAVE = rng('Bloco'), rng('Extensão (m)'), rng('Ativo de referência'), \
        rng('Poste fim'), rng('Chave do caminho (uso das fórmulas)')

    # --- Blocos (ativos um por coluna + concatenado por fórmula)
    wsb = wb.create_sheet('Blocos')
    HB = (['Bloco', 'Zona', 'Tronco/Ramal', 'Ramal', 'Ativo de referência', 'De']
          + [f'Próximo ativo {i}' for i in range(1, MAXN + 1)] + [f'Nível {i}' for i in range(1, NV + 1)]
          + ['Qtd. vãos', 'Extensão do bloco (m)', 'Qtd. ativos posteriores', 'Ativos posteriores (concatenado)',
             'Chaves 88 no bloco'] + [f'Posterior {i}' for i in range(1, MAXP + 1)])
    cb = {h: i + 1 for i, h in enumerate(HB)}
    rb = [[b['n'], f"T{b['zona']}", 'Tronco' if b['tronco'] else 'Ramal', b['ramal'], b['ref'], b['de'],
           *(b['ate'] + [None] * (MAXN - len(b['ate']))), *(b['caminho'] + [None] * (NV - len(b['caminho']))),
           None, None, None, None, ', '.join(b['chaves88']) or '—',
           *(b['posteriores'] + [None] * (MAXP - len(b['posteriores'])))] for b in blocos]
    tabela(wsb, 'Blocos', HB, rb, [7, 7, 12, 17, 12, 14] + [12] * MAXN + [8] * NV + [9, 12, 11, 42, 18] + [10] * MAXP)
    p1, p2 = L(cb['Posterior 1']), L(cb[f'Posterior {MAXP}'])
    for i, b in enumerate(blocos, 2):
        faixa = f'{p1}{i}:{p2}{i}'
        wsb.cell(i, cb['Qtd. vãos'], f'=COUNTIFS({R_BLOCO},A{i})')
        c = wsb.cell(i, cb['Extensão do bloco (m)'], f'=SUMIFS({R_EXT},{R_BLOCO},A{i})')
        c.number_format = '#,##0'
        wsb.cell(i, cb['Qtd. ativos posteriores'], f'=COUNTA({faixa})')
        wsb.cell(i, cb['Ativos posteriores (concatenado)'], f'=IF(COUNTA({faixa})=0,"—",_xlfn.TEXTJOIN(", ",TRUE,{faixa}))')
        for h in ('Qtd. vãos', 'Extensão do bloco (m)', 'Qtd. ativos posteriores', 'Ativos posteriores (concatenado)'):
            wsb.cell(i, cb[h]).font = body
        wsb.cell(i, 2).fill = zfill(b['zona'])
    tr = len(blocos) + 2
    wsb.cell(tr, 1, 'TOTAL').font = bold
    for h in ('Qtd. vãos', 'Extensão do bloco (m)'):
        c = wsb.cell(tr, cb[h], f'=SUM({L(cb[h])}2:{L(cb[h])}{tr - 1})')
        c.font, c.number_format = bold, '#,##0'
    wsb.cell(1, cb['Ativos posteriores (concatenado)']).comment = Comment(
        'Junta as colunas Posterior 1…N. Edite as colunas e o texto se atualiza.', 'Claude')
    wsb.freeze_panes = 'F2'
    CAT = lambda linha_bloco: f"=Blocos!{L(cb['Ativos posteriores (concatenado)'])}{linha_bloco}"
    linha_bloco = {b['n']: i for i, b in enumerate(blocos, 2)}

    # --- Por bloco (vãos agrupados; botão +/− na linha do bloco)
    wg = wb.create_sheet('Por bloco')
    HG = ['Bloco', 'Zona', 'Tronco/Ramal', 'Ramal', 'Ativo de referência', 'Trecho', 'Poste início', 'Poste fim',
          'Extensão (m)', 'Distância do ativo de referência até o fim (m)', 'Chave 88 no fim do vão',
          'Próximos ativos', 'Chaves 88 no bloco', 'Ativos posteriores']
    cabecalho(wg, HG)
    wg.sheet_properties.outlinePr.summaryBelow = False
    row = 2
    for b in blocos:
        vals = [b['n'], f"T{b['zona']}", 'Tronco' if b['tronco'] else 'Ramal', b['ramal'], b['ref'],
                f'=COUNTIFS({R_BLOCO},A{row})&" vãos"', None, None, f'=SUMIFS({R_EXT},{R_BLOCO},A{row})', None, None,
                ' / '.join(b['ate']), ', '.join(b['chaves88']) or '—', CAT(linha_bloco[b['n']])]
        for j, v in enumerate(vals, 1):
            c = wg.cell(row, j, v)
            c.font, c.fill = bold, zfill(b['zona'])
        wg.cell(row, 9).number_format = '#,##0'
        wg.row_dimensions[row].collapsed = True
        ini = row + 1
        for v in (x for x in vaos if x['bloco'] == b['n']):
            row += 1
            for j, val in ((6, v['id']), (7, v['ini']), (8, v['fim']), (9, v['ext']), (10, v['dist_ref']), (11, v['chave88_fim'] or None)):
                wg.cell(row, j, val).font = body
            wg.cell(row, 9).number_format = '0.0'
            wg.cell(row, 10).number_format = '#,##0'
        wg.row_dimensions.group(ini, row, outline_level=1, hidden=True)
        row += 1
    for i, w in enumerate([7, 7, 12, 17, 12, 10, 11, 11, 11, 16, 14, 24, 18, 42], 1):
        wg.column_dimensions[L(i)].width = w
    wg.freeze_panes = 'A2'

    # --- Árvore de blocos (visual da sequência, agrupado por tronco/ramal)
    wa = wb.create_sheet('Árvore de blocos')
    wa['A1'] = f'Árvore de blocos — {nome}'
    wa['A1'].font = Font(name=F, bold=True, size=14)
    wa['A2'] = ('Cada linha é um bloco. A célula colorida é o ativo de referência do bloco (cor da zona). '
                'As células em cinza à esquerda são os ativos antes dele, desde o DJ.')
    wa['A2'].font = Font(name=F, size=10, italic=True)
    wa['A3'] = 'Zonas:'
    wa['A3'].font = bold
    for z in range(1, max(b['zona'] for b in blocos) + 1):
        c = wa.cell(3, 1 + z, f'T{z}')
        c.fill, c.font, c.alignment = zfill(z), bold, Alignment(horizontal='center')
    HA = ['Bloco', 'Zona'] + [f'Nível {i}' for i in range(1, NV + 1)] + \
         ['De', 'Até (próximos ativos)', 'Vãos', 'Extensão (m)', 'Chaves 88 no bloco', 'Ativos englobados (depois do bloco)']
    cabecalho(wa, HA, 5)
    caixa = Border(*(Side(style='thin', color='2B3136'),) * 4)
    grupo = PatternFill('solid', fgColor='3A4248')
    titulo = lambda r: ('TRONCO — do DJ até o 2º religador' if r == 'Tronco' else
                        f'{r.upper()} — tudo depois do 2º religador' if r.startswith('Ramal após') else
                        f'{r.upper()} — sai do tronco no poste {r.split()[-1]}')
    row = 6
    for ram in dict.fromkeys(b['ramal'] for b in blocos):
        wa.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(HA))
        c = wa.cell(row, 1, titulo(ram))
        c.font, c.fill = Font(name=F, bold=True, size=10, color='FFFFFF'), grupo
        row += 1
        for b in (x for x in blocos if x['ramal'] == ram):
            zf = zfill(b['zona'])
            wa.cell(row, 1, b['n']).font = bold
            c = wa.cell(row, 2, f"T{b['zona']}")
            c.fill, c.font, c.alignment = zf, bold, Alignment(horizontal='center')
            for k_, a_ in enumerate(b['caminho']):
                c = wa.cell(row, 3 + k_, a_)
                c.alignment = Alignment(horizontal='center')
                if k_ == len(b['caminho']) - 1:
                    c.fill, c.font, c.border = zf, bold, caixa
                else:
                    c.font = Font(name=F, size=10, color='8A9096')
            base = 3 + NV
            wa.cell(row, base, b['de']).font = body
            wa.cell(row, base + 1, ' / '.join(b['ate'])).font = body
            wa.cell(row, base + 2, f'=COUNTIFS({R_BLOCO},A{row})').font = body
            c = wa.cell(row, base + 3, f'=SUMIFS({R_EXT},{R_BLOCO},A{row})')
            c.font, c.number_format = body, '#,##0'
            wa.cell(row, base + 4, ', '.join(b['chaves88']) or '—').font = body
            wa.cell(row, base + 5, CAT(linha_bloco[b['n']])).font = body
            for j in range(1, len(HA) + 1):
                if wa.cell(row, j).border.left.style is None:
                    wa.cell(row, j).border = Border(bottom=thin)
            row += 1
    for i, w in enumerate([7, 7] + [10] * NV + [13, 22, 7, 12, 18, 44], 1):
        wa.column_dimensions[L(i)].width = w
    wa.freeze_panes = 'C6'

    # --- Ativos
    wx = wb.create_sheet('Ativos')
    HX = ['Ativo', 'Tipo', 'Poste', 'No tronco', 'Zona antes', 'Zona depois', 'Bloco onde está',
          'Blocos em que é referência', 'Caminho desde o DJ', 'Vãos no próprio bloco', 'Extensão do próprio bloco (m)',
          'Vãos englobados (tudo a jusante)', 'Extensão englobada (m)', 'Ativos posteriores englobados',
          'Latitude', 'Longitude', 'Mapa']
    rx = [[a['nome'], TIPO[a['tipo']], a['poste'], 'sim' if a['no_tronco'] else 'não',
           f"T{a['zona_antes']}" if a['zona_antes'] else '—', f"T{a['zona_depois']}", a['bloco'] or '—',
           ', '.join(map(str, a['blocos_proprios'])) or '—', ' › '.join(a['caminho']), None, None, None, None,
           ', '.join(a['posteriores']) or '—', round(a['lat'], 7), round(a['lon'], 7), 'abrir'] for a in ativos]
    tabela(wx, 'Ativos', HX, rx, [10, 18, 9, 9, 10, 11, 10, 14, 30, 11, 14, 14, 14, 40, 12, 12, 8])
    for i, a in enumerate(ativos, 2):
        if a['tipo'] != '88':
            wx.cell(i, 10, f'=COUNTIFS({R_REF},A{i})').font = body
            c = wx.cell(i, 11, f'=SUMIFS({R_EXT},{R_REF},A{i})')
            c.font, c.number_format = body, '#,##0'
            wx.cell(i, 12, f'=SUMPRODUCT(--ISNUMBER(SEARCH("|"&A{i}&"|",{R_CHAVE})))').font = body
            c = wx.cell(i, 13, f'=SUMPRODUCT(ISNUMBER(SEARCH("|"&A{i}&"|",{R_CHAVE}))*{R_EXT})')
            c.font, c.number_format = body, '#,##0'
        c = wx.cell(i, 17)
        c.hyperlink, c.font = f"https://www.google.com/maps?q={a['lat']:.7f},{a['lon']:.7f}", link
    wx.cell(1, 12).comment = Comment('Vãos depois deste ativo, em qualquer bloco. Vazio para chaves 88, que não '
                                     'são ativo de referência.', 'Claude')
    wx.freeze_panes = 'B2'

    # --- Consulta
    wc = wb.create_sheet('Consulta')
    wc['A1'] = 'Onde fica este poste?'
    wc['A1'].font = Font(name=F, bold=True, size=14)
    wc['A2'] = 'Digite o código do poste na célula amarela (o poste no fim do vão). A planilha mostra o vão e a sequência desde o DJ.'
    wc['A2'].font = Font(name=F, size=10, italic=True)
    wc['A4'] = 'Poste'
    ex = next((v for v in vaos if v['id'] == defeito), vaos[-1])
    wc['B4'] = ex['fim']
    wc['B4'].fill, wc['B4'].font = PatternFill('solid', fgColor='FFFF00'), Font(name=F, bold=True, size=11)
    wc['C4'] = '← exemplo'
    wc['C4'].font = Font(name=F, size=9, italic=True, color='666666')
    consulta = [('Trecho (vão)', 'Trecho'), ('Poste início', 'Poste início'), ('Zona', 'Zona'),
                ('Tronco/Ramal', 'Tronco/Ramal'), ('Ramal', 'Ramal'), ('Bloco', 'Bloco'),
                ('Ativo de referência', 'Ativo de referência'),
                ('Distância do ativo de referência (m)', 'Distância do ativo de referência até o fim (m)'),
                ('Distância do DJ (m)', 'Distância do DJ até o fim (m)'),
                ('Ativos posteriores englobados', 'Ativos posteriores englobados'),
                ('Chave 88 no fim do vão', 'Chave 88 no fim do vão'), ('Caminho desde o DJ', 'Caminho desde o DJ'),
                ('Latitude', 'Latitude (meio do vão)'), ('Longitude', 'Longitude (meio do vão)')]
    for j, (lab, h) in enumerate(consulta, 6):
        wc.cell(j, 1, lab).font = bold
        c = wc.cell(j, 2, f'=IFERROR(INDEX({rng(h)},MATCH($B$4,{R_POSTE},0)),"poste não encontrado")')
        c.font = body
        if 'Distância' in lab:
            c.number_format = '#,##0'
    wc.column_dimensions['A'].width, wc.column_dimensions['B'].width, wc.column_dimensions['C'].width = 36, 46, 20

    wb._sheets = [wb[n] for n in ['Leia-me', 'Árvore de blocos', 'Consulta', 'Blocos', 'Por bloco', 'Trechos', 'Ativos']]
    wb.active = 0
    wb.save(caminho)


# ============================================================ MAPA
def dados_mapa(res, defeito=None):
    prox = _prox(res)
    return dict(
        trechos=[dict(id=v['id'], a=[round(x, 7) for x in v['a']], b=[round(x, 7) for x in v['b']], ini=v['ini'],
                      fim=v['fim'], ext=round(v['ext'], 3), dj=round(v['dist_dj']), dref=round(v['dist_ref']),
                      z=v['zona'], tr='Tronco' if v['tronco'] else 'Ramal', ramal=v['ramal'], bloco=v['bloco'],
                      ref=v['ref'], path=v['caminho'], prox=' / '.join(prox[v['bloco']]), post=v['posteriores'],
                      c88=v['chave88_fim']) for v in res['vaos']],
        ativos=[dict(nome=a['nome'], t=a['tipo'], lat=round(a['lat'], 7), lon=round(a['lon'], 7), zin=a['zona_antes'],
                     zout=a['zona_depois'], poste=a['poste'], post=a['posteriores'], path=a['caminho'], bloco=a['bloco'])
                for a in res['ativos']],
        blocos=[dict(n=b['n'], z=b['zona'], r=b['ramal'], ref=b['ref'], de=b['de'], ate=b['ate'], path=b['caminho'],
                     post=b['posteriores'], c88=b['chaves88'], qtd=b['qtd'], ext=round(b['ext']),
                     pos=[round(x, 7) for x in b['pos']]) for b in res['blocos']],
        defeito=defeito, avisos=res['avisos'])


def mapa(res, nome, caminho, defeito=None, aviso_exemplo=False):
    modelo = (Path(__file__).resolve().parent / 'mapa_modelo.html').read_text(encoding='utf-8')
    rotulo = 'Circuito de exemplo · não é rede real' if aviso_exemplo else f'Alimentador {nome}'
    html = (modelo.replace('%%TITULO%%', escape(nome)).replace('%%ROTULO%%', escape(rotulo))
            .replace('%%DADOS%%', json.dumps(dados_mapa(res, defeito), ensure_ascii=False, separators=(',', ':'))
                     .replace('</', '<\\/')))
    assert '%%' not in html
    Path(caminho).write_text(html, encoding='utf-8')
